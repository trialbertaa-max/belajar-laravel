import os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import lxml.etree as ET

xslt_path = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
xslt = ET.parse(xslt_path)
transform = ET.XSLT(xslt)

def mml_to_omml_element(mml_str):
    mml_tree = ET.fromstring(mml_str)
    omml_tree = transform(mml_tree)
    root = omml_tree.getroot()
    if root.tag.endswith('oMath'):
        para = ET.Element('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara')
        para_pr = ET.SubElement(para, '{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathParaPr')
        jc = ET.SubElement(para_pr, '{http://schemas.openxmlformats.org/officeDocument/2006/math}jc')
        jc.set('{http://schemas.openxmlformats.org/officeDocument/2006/math}val', 'center')
        para.append(root)
        return para
    return root

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, size_pt=12):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt)
    return r

def main():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.18)    # 3 cm
        s.bottom_margin = Inches(1.18) # 3 cm
        s.left_margin = Inches(1.57)   # 4 cm
        s.right_margin = Inches(1.18)  # 3 cm

    # Heading 3.7
    h = doc.add_paragraph()
    format_p(h, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h, '3.7\tPemodelan Derau Sintetis', bold=True)

    # Paragraf 1
    p1 = doc.add_paragraph()
    format_p(p1)
    add_styled_run(p1, 'Degradasi citra disimulasikan menggunakan model derau campuran Poisson-Gaussian terkontrol untuk merepresentasikan karakteristik fisis derau pada radiografi digital proyeksi dosis rendah (Foi et al., 2008). Model ini menggabungkan fluktuasi foton kuantum sinar-X yang bergantung pada intensitas dengan derau termal elektronik detektor yang independen dari intensitas. Formulasi matematis pembentukan citra terdegradasi dinyatakan sebagai berikut:')

    # Persamaan (Tanpa nomor persamaan, rata tengah persis seperti di Bab II)
    mml_eq = '''<math xmlns="http://www.w3.org/1998/Math/MathML">
      <mrow>
        <mi>z</mi><mo>(</mo><mi>x</mi><mo>,</mo><mi>y</mi><mo>)</mo>
        <mo>=</mo>
        <mi>f</mi><mo>(</mo><mi>x</mi><mo>,</mo><mi>y</mi><mo>)</mo>
        <mo>+</mo>
        <msub><mi>&#951;</mi><mi>p</mi></msub><mo>(</mo><mi>f</mi><mo>(</mo><mi>x</mi><mo>,</mo><mi>y</mi><mo>)</mo><mo>)</mo>
        <mo>+</mo>
        <msub><mi>&#951;</mi><mi>g</mi></msub><mo>(</mo><mi>x</mi><mo>,</mo><mi>y</mi><mo>)</mo>
      </mrow>
    </math>'''

    p_eq = doc.add_paragraph()
    format_p(p_eq, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, space_before_pt=4, line_spacing=1.0, first_line_indent_pt=0)
    omml_elm = mml_to_omml_element(mml_eq)
    p_eq._p.append(omml_elm)

    # Paragraf penjelasan parameter
    p2 = doc.add_paragraph()
    format_p(p2)
    add_styled_run(p2, 'dengan ')
    add_styled_run(p2, 'z(x, y)', italic=True)
    add_styled_run(p2, ' menyatakan intensitas piksel citra terdegradasi pada koordinat spasial ')
    add_styled_run(p2, '(x, y)', italic=True)
    add_styled_run(p2, ', ')
    add_styled_run(p2, 'f(x, y)', italic=True)
    add_styled_run(p2, ' merepresentasikan intensitas citra acuan bebas derau (')
    add_styled_run(p2, 'clean reference', italic=True)
    add_styled_run(p2, '), ')
    add_styled_run(p2, '\u03b7p(f(x, y))', italic=True)
    add_styled_run(p2, ' memodelkan komponen derau Poisson akibat fluktuasi kedatangan foton kuantum (')
    add_styled_run(p2, 'quantum mottle', italic=True)
    add_styled_run(p2, '), dan ')
    add_styled_run(p2, '\u03b7g(x, y)', italic=True)
    add_styled_run(p2, ' memodelkan komponen derau aditif Gaussian N(0, \u03c3n\u00b2) dari sirkuit instrumen detektor.')

    # Paragraf tingkat derau
    p3 = doc.add_paragraph()
    format_p(p3)
    add_styled_run(p3, 'Pengujian ketahanan restorasi dilakukan pada tiga tingkatan deviasi standar derau Gaussian terkontrol, yaitu:')

    levels = [
        ('1.', 'Derau rendah (\u03c3n = 15): menyimulasikan degradasi akibat penurunan fluks foton ringan.'),
        ('2.', 'Derau sedang (\u03c3n = 25): menyimulasikan kondisi radiografi toraks dosis rendah standar.'),
        ('3.', 'Derau berat (\u03c3n = 50): menyimulasikan kondisi radiografi toraks dosis sangat rendah (ultra-low-dose).')
    ]

    for num, text in levels:
        p_lvl = doc.add_paragraph()
        format_p(p_lvl, space_after_pt=1, space_before_pt=0, first_line_indent_pt=14.17)
        add_styled_run(p_lvl, num + '\t')
        words = text.split(' ')
        for w_idx, w in enumerate(words):
            sp = ' ' if w_idx < len(words) - 1 else ''
            clean_w = w.strip('(),:.\u03c3')
            is_it = clean_w in ['ultra-low-dose', 'n', '\u03c3n']
            add_styled_run(p_lvl, w + sp, italic=is_it)

    # Paragraf batasan masalah (BERHENTI DI SINI TANPA KALIMAT PENGANTAR GAMBAR)
    p4 = doc.add_paragraph()
    format_p(p4)
    add_styled_run(p4, 'Sesuai batasan masalah penelitian, penambahan derau sintetis ini murni difungsikan sebagai simulasi degradasi derau terkontrol untuk mengevaluasi efektivitas filter restorasi dan dampaknya terhadap biometri CTR, serta tidak digunakan untuk mengestimasi dosis radiasi fisik aktual pasien.')

    # Gambar 3.2 langsung diletakkan di bawah paragraf
    img_path = r'd:\laragon\www\belajar-laravel\Gambar_3_2_Simulasi_Derau_Poisson_Gaussian.png'
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        format_p(p_img, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, space_before_pt=8, line_spacing=1.0, first_line_indent_pt=0)
        p_img.add_run().add_picture(img_path, width=Inches(5.4))

        # Caption Gambar 3.2
        p_cap = doc.add_paragraph()
        format_p(p_cap, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=6, space_before_pt=2, line_spacing=1.15, first_line_indent_pt=0)
        add_styled_run(p_cap, 'Gambar 3.2 Simulasi Degradasi Derau Terkontrol (Poisson-Gaussian) pada Citra Radiografi Toraks PA', bold=False)

    out_file = r'd:\laragon\www\belajar-laravel\SUBBAB_3_7_PEMODELAN_DERAU_FINAL.docx'
    doc.save(out_file)
    print('Successfully generated final clean Subbab 3.7:', out_file)

if __name__ == '__main__':
    main()
