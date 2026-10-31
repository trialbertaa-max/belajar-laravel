import docx
import re

doc = docx.Document('DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx')

foreign_words = [
    'denoising', 'convolutional', 'neural', 'network', 'residual', 'learning',
    'batch', 'normalization', 'deep', 'cnn', 'vanishing', 'gradient', 'loss',
    'function', 'mean', 'squared', 'error', 'supervised', 'patch', 'patches',
    'ground', 'truth', 'pooling', 'fully-connected', 'downsampling', 'filter',
    'filters', 'feature', 'maps', 'internal', 'covariate', 'shift', 'receptive',
    'field', 'kernel', 'natural', 'images', 'additive', 'white', 'noise',
    'hallucination', 'risk', 'domain', 'adaptation', 'transfer', 'gaussian',
    'rectified', 'linear', 'unit', 'layer', 'layers'
]

# Find bounds of 2.4.3
start_idx = None
end_idx = None
for i, p in enumerate(doc.paragraphs):
    if '2.4.3' in p.text:
        start_idx = i
    if '2.4.4' in p.text:
        end_idx = i
        break

print(f"=== Scanning 2.4.3 (p[{start_idx}] to p[{end_idx}]) in DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx ===")
flagged = 0
for p_idx in range(start_idx, end_idx):
    p = doc.paragraphs[p_idx]
    for r_idx, r in enumerate(p.runs):
        if not r.italic:
            text_lower = r.text.lower()
            for fw in foreign_words:
                # check word boundaries
                matches = re.findall(rf'\b{fw}\b', text_lower)
                if matches:
                    # Note: 'internal' in Indonesian is an accepted KBBI word, but let's check
                    print(f"P{p_idx} Run[{r_idx}] '{r.text}' contains non-italic '{fw}'")
                    flagged += 1

print(f"=== Scan completed: {flagged} potential non-italic foreign words found ===")
