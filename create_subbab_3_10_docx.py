import os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False, subscript=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic
    r.font.subscript = subscript

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, size_pt=12, subscript=False):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt, subscript=subscript)
    return r

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def style_apa_table(table, header_bg='F8FAFC'):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    num_rows = len(table.rows)
    for r_idx, row in enumerate(table.rows):
        # Prevent row split across pages
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if r_idx == 0:
            trPr.append(OxmlElement('w:tblHeader'))

        for cell in row.cells:
            # Set top/bottom borders per APA standard
            if r_idx == 0:
                set_cell_border(cell, top={'val': 'single', 'sz': 8, 'color': '000000'},
                                      bottom={'val': 'single', 'sz': 6, 'color': '000000'},
                                      left={'val': 'none'}, right={'val': 'none'})
            elif r_idx == num_rows - 1:
                set_cell_border(cell, top={'val': 'single', 'sz': 4, 'color': 'CCCCCC'},
                                      bottom={'val': 'single', 'sz': 8, 'color': '000000'},
                                      left={'val': 'none'}, right={'val': 'none'})
            else:
                set_cell_border(cell, top={'val': 'none'}, bottom={'val': 'none'},
                                      left={'val': 'none'}, right={'val': 'none'})

def main():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.18)    # 3 cm
        s.bottom_margin = Inches(1.18) # 3 cm
        s.left_margin = Inches(1.57)   # 4 cm
        s.right_margin = Inches(1.18)  # 3 cm

    # Heading 3.10
    h = doc.add_paragraph()
    format_p(h, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h, '3.10\tEvaluasi Kualitas Citra dan Pelestarian Batas Anatomi', bold=True)

    # Paragraf Pengantar
    p1 = doc.add_paragraph()
    format_p(p1)
    add_styled_run(p1, 'Evaluasi performa pemrosesan citra dilakukan melalui dua tingkatan pengujian kuantitatif, yaitu fidelitas citra secara objektif dan pengujian berbasis tugas klinis (')
    add_styled_run(p1, 'task-based evaluation', italic=True)
    add_styled_run(p1, '):')

    # Poin 1: Evaluasi Kualitas Citra Objektif
    p_eq = doc.add_paragraph()
    format_p(p_eq, space_after_pt=2, space_before_pt=0, first_line_indent_pt=14.17)
    add_styled_run(p_eq, '1.\tEvaluasi Kualitas Citra Objektif: Menilai tingkat perbaikan kualitas visual berbasis piksel terhadap citra acuan bebas derau (')
    add_styled_run(p_eq, 'clean reference', italic=True)
    add_styled_run(p_eq, ') menggunakan metrik ')
    add_styled_run(p_eq, 'Peak Signal-to-Noise Ratio', italic=True)
    add_styled_run(p_eq, ' (PSNR) dalam satuan desibel (dB) dan ')
    add_styled_run(p_eq, 'Structural Similarity Index Measure', italic=True)
    add_styled_run(p_eq, ' (SSIM) untuk mengukur kemiripan struktural perseptual luminansi, kontras, dan struktur.')

    # Judul Tabel 3.1
    p_t1_title = doc.add_paragraph()
    format_p(p_t1_title, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=3, space_before_pt=6, line_spacing=1.15, first_line_indent_pt=0)
    add_styled_run(p_t1_title, 'Tabel 3.1 Rancangan Matriks Hasil Evaluasi Kualitas Citra Pasca Restorasi', bold=False)

    # Tabel 3.1
    t1_data = [
        ['No Citra', 'Tingkat Derau (\u03c3n)', 'PSNR Gaussian (dB)', 'PSNR Bilateral (dB)', 'PSNR DnCNN (dB)', 'SSIM Gaussian', 'SSIM Bilateral', 'SSIM DnCNN'],
        ['1', '\u03c3n = 15', '...', '...', '...', '...', '...', '...'],
        ['2', '\u03c3n = 25', '...', '...', '...', '...', '...', '...'],
        ['3', '\u03c3n = 50', '...', '...', '...', '...', '...', '...'],
        ['...', '...', '...', '...', '...', '...', '...', '...'],
        ['Rerata', '-', '...', '...', '...', '...', '...', '...']
    ]

    t1 = doc.add_table(rows=len(t1_data), cols=len(t1_data[0]))
    style_apa_table(t1)
    col_widths_1 = [Inches(0.6), Inches(0.8), Inches(0.7), Inches(0.7), Inches(0.7), Inches(0.65), Inches(0.65), Inches(0.65)]
    for r_idx, row_data in enumerate(t1_data):
        for c_idx, val in enumerate(row_data):
            cell = t1.cell(r_idx, c_idx)
            cell.width = col_widths_1[c_idx]
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER
            format_p(p, align=align, space_after_pt=2, space_before_pt=2, line_spacing=1.0, first_line_indent_pt=0)
            is_b = (r_idx == 0 or r_idx == len(t1_data) - 1)
            add_styled_run(p, val, bold=is_b, size_pt=10.5)

    # Spasi setelah tabel 1
    p_sp1 = doc.add_paragraph()
    format_p(p_sp1, space_after_pt=4, space_before_pt=0, line_spacing=1.0, first_line_indent_pt=0)

    # Poin 2: Evaluasi Batas Anatomi & Biometri CTR
    p_eb = doc.add_paragraph()
    format_p(p_eb, space_after_pt=2, space_before_pt=0, first_line_indent_pt=14.17)
    add_styled_run(p_eb, '2.\tEvaluasi Pelestarian Batas Anatomi dan Akurasi Biometrik CTR: Untuk memastikan reduksi derau tidak memicu pengaburan tepi organ (')
    add_styled_run(p_eb, 'over-smoothing', italic=True)
    add_styled_run(p_eb, ') pada Bilateral Filter atau distorsi halusinasi pada DnCNN, evaluasi klinis mencakup:')

    # Sub-butir a dan b
    p_sub_a = doc.add_paragraph()
    format_p(p_sub_a, space_after_pt=1, space_before_pt=0, first_line_indent_pt=28.35)
    add_styled_run(p_sub_a, 'a.\tPelestarian batas anatomi (')
    add_styled_run(p_sub_a, 'boundary preservation', italic=True)
    add_styled_run(p_sub_a, '): Mengukur derajat kesesuaian kontur tepi siluet jantung dan sangkar toraks menggunakan ')
    add_styled_run(p_sub_a, 'Dice Similarity Coefficient', italic=True)
    add_styled_run(p_sub_a, ' (DSC) untuk tumpang-tindih area, ')
    add_styled_run(p_sub_a, '95th Percentile Hausdorff Distance', italic=True)
    add_styled_run(p_sub_a, ' (HD95) untuk jarak deviasi batas maksimum terluar, dan ')
    add_styled_run(p_sub_a, 'Mean Surface Distance', italic=True)
    add_styled_run(p_sub_a, ' (MSD) untuk simpangan rata-rata jarak permukaan batas organ.')

    p_sub_b = doc.add_paragraph()
    format_p(p_sub_b, space_after_pt=2, space_before_pt=0, first_line_indent_pt=28.35)
    add_styled_run(p_sub_b, 'b.\tAkurasi pengukuran CTR: Menghitung selisih dan galat biometri terhadap standar referensi penelitian menggunakan ')
    add_styled_run(p_sub_b, 'Mean Absolute Error', italic=True)
    add_styled_run(p_sub_b, ' (MAE), deviasi rata-rata (')
    add_styled_run(p_sub_b, 'Bias', italic=True)
    add_styled_run(p_sub_b, '), serta penetapan batas kesepakatan klinis ')
    add_styled_run(p_sub_b, 'Limits of Agreement', italic=True)
    add_styled_run(p_sub_b, ' (LoA) melalui analisis plot Bland-Altman.')

    # Judul Tabel 3.2
    p_t2_title = doc.add_paragraph()
    format_p(p_t2_title, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=3, space_before_pt=6, line_spacing=1.15, first_line_indent_pt=0)
    add_styled_run(p_t2_title, 'Tabel 3.2 Rancangan Matriks Evaluasi Akurasi Pengukuran CTR Antar Metode', bold=False)

    # Tabel 3.2
    t2_data = [
        ['Metode Perlakuan', 'Rerata CTR (%)', 'MAE CTR (%)', 'Bias (%)', 'DSC', 'HD95 (mm)', 'MSD (mm)', 'Bland-Altman LoA (%)'],
        ['Citra Berderau (Noisy Baseline)', '...', '...', '...', '...', '...', '...', '...'],
        ['Gaussian Filter Baseline', '...', '...', '...', '...', '...', '...', '...'],
        ['Bilateral Filter', '...', '...', '...', '...', '...', '...', '...'],
        ['DnCNN', '...', '...', '...', '...', '...', '...', '...']
    ]

    t2 = doc.add_table(rows=len(t2_data), cols=len(t2_data[0]))
    style_apa_table(t2)
    col_widths_2 = [Inches(1.5), Inches(0.55), Inches(0.55), Inches(0.5), Inches(0.45), Inches(0.55), Inches(0.55), Inches(0.85)]
    for r_idx, row_data in enumerate(t2_data):
        for c_idx, val in enumerate(row_data):
            cell = t2.cell(r_idx, c_idx)
            cell.width = col_widths_2[c_idx]
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            format_p(p, align=align, space_after_pt=2, space_before_pt=2, line_spacing=1.0, first_line_indent_pt=0)
            is_b = (r_idx == 0)
            add_styled_run(p, val, bold=is_b, size_pt=10.5)

    out_file = r'd:\laragon\www\belajar-laravel\SUBBAB_3_10_EVALUASI_KUALITAS.docx'
    doc.save(out_file)
    print('Successfully generated Subbab 3.10:', out_file)

if __name__ == '__main__':
    main()

