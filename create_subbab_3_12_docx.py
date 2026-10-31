import os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False, subscript=False, superscript=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic
    r.font.subscript = subscript
    r.font.superscript = superscript

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, size_pt=12, subscript=False, superscript=False):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt, subscript=subscript, superscript=superscript)
    return r

def main():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.18)    # 3 cm
        s.bottom_margin = Inches(1.18) # 3 cm
        s.left_margin = Inches(1.57)   # 4 cm
        s.right_margin = Inches(1.18)  # 3 cm

    # Heading 3.12
    h = doc.add_paragraph()
    format_p(h, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h, '3.12\tAnalisis Statistik Inferensial', bold=True)

    # Paragraf Pengantar
    p1 = doc.add_paragraph()
    format_p(p1)
    add_styled_run(p1, 'Uji statistik inferensial dirancang untuk menguji signifikansi perbedaan performa antar-metode restorasi citra. Variabel keluaran penelitian dikelompokkan menjadi luaran primer (')
    add_styled_run(p1, 'primary outcome', italic=True)
    add_styled_run(p1, ') berupa galat biometri (MAE CTR) serta luaran sekunder (')
    add_styled_run(p1, 'secondary outcomes', italic=True)
    add_styled_run(p1, ') yang mencakup metrik kualitas citra (PSNR dan SSIM), galat batas anatomi (DSC, HD95, MSD), dan waktu komputasi (')
    add_styled_run(p1, 'runtime', italic=True)
    add_styled_run(p1, '). Prosedur analisis statistik dieksekusi melalui tahapan berikut:')

    # Poin 1-4
    points = [
        ("1.", "Uji Normalitas: Distribusi data galat pengukuran diuji menggunakan uji Shapiro-Wilk pada tingkat signifikansi \u03b1 = 0,05."),
        ("2.", "Uji Komparasi Multikelompok Berpasangan: Karena setiap perlakuan dievaluasi pada subjek citra yang sama (within-subject design), komparasi antar-kondisi (clean reference, noisy baseline, Gaussian, Bilateral Filter, dan DnCNN) dianalisis menggunakan Repeated Measures ANOVA jika data berdistribusi normal, atau uji non-parametrik Friedman jika asumsi normalitas tidak terpenuhi. Besaran efek (effect size) dilaporkan menggunakan partial eta-squared (\u03b7p\u00b2) atau Kendall's W pada rentang interval kepercayaan 95% (95% CI)."),
        ("3.", "Uji Lanjut Berpasangan (Post-Hoc): Apabila uji komparasi multikelompok menunjukkan perbedaan yang signifikan (p < 0,05), analisis perbandingan berpasangan dilanjutkan menggunakan paired t-test (untuk data normal) atau Wilcoxon signed-rank test (untuk data non-parametrik) dengan penyesuaian nilai p menggunakan koreksi Holm (Holm, 1979) guna mengendalikan tingkat kesalahan tipe I (family-wise error rate)."),
        ("4.", "Analisis Kesepakatan Klinis (Agreement Analysis): Evaluasi kesepakatan pengukuran biometri CTR antar-metode terhadap standar referensi dinilai menggunakan metode Bland-Altman untuk menentukan deviasi rata-rata (Bias) dan Limits of Agreement (LoA). Analisis korelasi Pearson atau Spearman digunakan sebagai pengujian pendukung untuk menilai linearitas asosiasi.")
    ]

    for num, text in points:
        p = doc.add_paragraph()
        format_p(p, space_after_pt=1, space_before_pt=0, first_line_indent_pt=14.17)
        add_styled_run(p, num + '\t')
        words = text.split(' ')
        for w_idx, w in enumerate(words):
            sp = ' ' if w_idx < len(words) - 1 else ''
            clean_w = w.strip('(),:.')
            is_it = clean_w in [
                'within-subject', 'design', 'clean', 'reference', 'noisy', 'baseline',
                'Repeated', 'Measures', 'ANOVA', 'Friedman', 'effect', 'size',
                'partial', 'eta-squared', 'Kendall\'s', 'W', 'CI', 'Post-Hoc',
                'paired', 't-test', 'Wilcoxon', 'signed-rank', 'test', 'p',
                'Holm', 'family-wise', 'error', 'rate', 'Agreement', 'Analysis',
                'Bland-Altman', 'Bias', 'Limits', 'of', 'Agreement', 'LoA',
                'Pearson', 'Spearman'
            ]
            add_styled_run(p, w + sp, italic=is_it)

    out_file = r'd:\laragon\www\belajar-laravel\SUBBAB_3_12_ANALISIS_STATISTIK.docx'
    doc.save(out_file)
    print('Generated 3.12:', out_file)

if __name__ == '__main__':
    main()

