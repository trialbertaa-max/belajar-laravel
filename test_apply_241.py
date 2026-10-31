import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.text.paragraph import Paragraph
import lxml.etree as ET
import shutil
import copy

# Load XSLT
xslt_path = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
xslt = ET.parse(xslt_path)
transform = ET.XSLT(xslt)

def mml_to_omml_element(mml_str):
    mml_tree = ET.fromstring(mml_str)
    omml_tree = transform(mml_tree)
    root = omml_tree.getroot()
    if root.tag.endswith('oMath'):
        para = ET.Element('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara')
        para_pr = ET.SubElement(para, '{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathParaPr')
        jc = ET.SubElement(para_pr, '{http://schemas.openxmlformats.org/officeDocument/2006/math}jc')
        jc.set('{http://schemas.openxmlformats.org/officeDocument/2006/math}val', 'center')
        para.append(root)
        root = para
    xml_bytes = ET.tostring(root, encoding='utf-8')
    return parse_xml(xml_bytes)

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False, subscript=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic
    if subscript:
        r.font.subscript = True

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, subscript=False, size_pt=12):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, subscript=subscript, size_pt=size_pt)
    return r

def build_section_241(target_before_p):
    def create_p(style_name='Normal'):
        new_p_elm = OxmlElement('w:p')
        target_before_p._p.addprevious(new_p_elm)
        p = Paragraph(new_p_elm, target_before_p._parent)
        if style_name != 'Normal':
            p.style = style_name
        return p

    def create_eq_p(mml_str):
        p = create_p()
        format_p(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, line_spacing=1.0, first_line_indent_pt=0)
        omml_elm = mml_to_omml_element(mml_str)
        p._p.append(omml_elm)
        return p

    # 1. Heading 3: 2.4.1 Gaussian Filter
    h3 = create_p('Heading 3')
    format_p(h3, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=4, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h3, "2.4.1\t", bold=True, italic=False)
    add_styled_run(h3, "Gaussian Filter", bold=True, italic=True)

    # 2. Paragraf 1: Definisi & Kernel Kontinu 2D
    p1 = create_p()
    format_p(p1)
    add_styled_run(p1, "Gaussian filter", italic=True)
    add_styled_run(p1, " merupakan metode penapisan linear dalam ranah spasial (")
    add_styled_run(p1, "spatial domain", italic=True)
    add_styled_run(p1, ") yang beroperasi melalui operasi konvolusi diskret dua dimensi menggunakan fungsi pembobotan terdistribusi normal (")
    add_styled_run(p1, "Gaussian distribution", italic=True)
    add_styled_run(p1, ") (Gonzalez & Woods, 2018). Fungsi pembobot ")
    add_styled_run(p1, "kernel Gaussian", italic=True)
    add_styled_run(p1, " kontinu dua dimensi didefinisikan melalui persamaan matematis:")

    # 3. Persamaan 1: Kernel Gaussian Kontinu
    mml1 = '''<math xmlns="http://www.w3.org/1998/Math/MathML">
      <msub><mi>G</mi><mi>σ</mi></msub>
      <mo>(</mo><mi>x</mi><mo>,</mo><mi>y</mi><mo>)</mo>
      <mo>=</mo>
      <mfrac>
        <mn>1</mn>
        <mrow><mn>2</mn><mi>π</mi><msup><mi>σ</mi><mn>2</mn></msup></mrow>
      </mfrac>
      <mi>exp</mi>
      <mfenced>
        <mrow>
          <mo>-</mo>
          <mfrac>
            <mrow><msup><mi>x</mi><mn>2</mn></msup><mo>+</mo><msup><mi>y</mi><mn>2</mn></msup></mrow>
            <mrow><mn>2</mn><msup><mi>σ</mi><mn>2</mn></msup></mrow>
          </mfrac>
        </mrow>
      </mfenced>
    </math>'''
    create_eq_p(mml1)

    # 4. Paragraf 2: Diskretisasi & Operasi Konvolusi
    p2 = create_p()
    format_p(p2)
    add_styled_run(p2, "Pada implementasi citra digital, operasi penyaringan pada piksel target berkoordinat ")
    add_styled_run(p2, "p", italic=True)
    add_styled_run(p2, " = (")
    add_styled_run(p2, "x", italic=True)
    add_styled_run(p2, ", ")
    add_styled_run(p2, "y", italic=True)
    add_styled_run(p2, ") dihitung melalui diskritisasi fungsi pembobot pada jendela lingkungan spasial lokal ")
    add_styled_run(p2, "Ω", italic=True)
    add_styled_run(p2, ". Nilai intensitas piksel terestorasi ")
    add_styled_run(p2, "I", italic=True)
    add_styled_run(p2, "Gaussian", italic=False, subscript=True)
    add_styled_run(p2, " (")
    add_styled_run(p2, "p", italic=True)
    add_styled_run(p2, ") diperoleh melalui penjumlahan terbobot dari intensitas piksel tetangga ")
    add_styled_run(p2, "q", italic=True)
    add_styled_run(p2, " \u2208 ")
    add_styled_run(p2, "Ω", italic=True)
    add_styled_run(p2, " yang dinormalisasi dengan faktor pembobot ")
    add_styled_run(p2, "W", italic=True)
    add_styled_run(p2, "p", italic=True, subscript=True)
    add_styled_run(p2, ":")

    # 5. Persamaan 2: Diskret Konvolusi
    mml2 = '''<math xmlns="http://www.w3.org/1998/Math/MathML">
      <msub><mi>I</mi><mtext>Gaussian</mtext></msub>
      <mo>(</mo><mi>p</mi><mo>)</mo>
      <mo>=</mo>
      <mfrac>
        <mn>1</mn>
        <msub><mi>W</mi><mi>p</mi></msub>
      </mfrac>
      <munder>
        <mo>∑</mo>
        <mrow><mi>q</mi><mo>∈</mo><mi>Ω</mi></mrow>
      </munder>
      <msub><mi>G</mi><mi>σ</mi></msub>
      <mo>(</mo><mo>‖</mo><mi>p</mi><mo>-</mo><mi>q</mi><mo>‖</mo><mo>)</mo>
      <mo>·</mo>
      <mi>I</mi>
      <mo>(</mo><mi>q</mi><mo>)</mo>
    </math>'''
    create_eq_p(mml2)

    # 6. Persamaan 3: Normalisasi W_p
    mml3 = '''<math xmlns="http://www.w3.org/1998/Math/MathML">
      <msub><mi>W</mi><mi>p</mi></msub>
      <mo>=</mo>
      <munder>
        <mo>∑</mo>
        <mrow><mi>q</mi><mo>∈</mo><mi>Ω</mi></mrow>
      </munder>
      <msub><mi>G</mi><mi>σ</mi></msub>
      <mo>(</mo><mo>‖</mo><mi>p</mi><mo>-</mo><mi>q</mi><mo>‖</mo><mo>)</mo>
    </math>'''
    create_eq_p(mml3)

    # 7. Paragraf 3: Keterangan Parameter & Mekanisme Peredaman Derau
    p3 = create_p()
    format_p(p3)
    add_styled_run(p3, "Keterangan parameter matematis: ")
    add_styled_run(p3, "Ω", italic=True)
    add_styled_run(p3, " merepresentasikan jendela lingkungan spasial lokal (")
    add_styled_run(p3, "spatial neighborhood window", italic=True)
    add_styled_run(p3, ") yang berpusat pada piksel target ")
    add_styled_run(p3, "p", italic=True)
    add_styled_run(p3, ", ||")
    add_styled_run(p3, "p", italic=True)
    add_styled_run(p3, " \u2212 ")
    add_styled_run(p3, "q", italic=True)
    add_styled_run(p3, "|| menyatakan jarak ")
    add_styled_run(p3, "Euclidean", italic=True)
    add_styled_run(p3, " antar-piksel dalam bidang citra, parameter ")
    add_styled_run(p3, "σ", italic=True)
    add_styled_run(p3, " merupakan simpangan baku spasial (")
    add_styled_run(p3, "spatial standard deviation", italic=True)
    add_styled_run(p3, ") yang mengatur lebar sebaran fungsi pembobot serta skala penghalusan (")
    add_styled_run(p3, "smoothing scale", italic=True)
    add_styled_run(p3, "), dan ")
    add_styled_run(p3, "W", italic=True)
    add_styled_run(p3, "p", italic=True, subscript=True)
    add_styled_run(p3, " adalah faktor penormalan untuk menjamin kelestarian rerata intensitas energi foton (")
    add_styled_run(p3, "energy conservation", italic=True)
    add_styled_run(p3, " atau ")
    add_styled_run(p3, "unit gain", italic=True)
    add_styled_run(p3, "). Secara fisis, ")
    add_styled_run(p3, "Gaussian filter", italic=True)
    add_styled_run(p3, " bekerja sangat efektif meredam fluktuasi derau acak frekuensi tinggi (")
    add_styled_run(p3, "high-frequency random noise", italic=True)
    add_styled_run(p3, ") pada area homogen melalui perataan rata-rata lokal (")
    add_styled_run(p3, "local averaging", italic=True)
    add_styled_run(p3, ").")

    # 8. Paragraf 4: Kelemahan Isotropik & Edge Blurring pada Toraks
    p4 = create_p()
    format_p(p4)
    add_styled_run(p4, "Namun demikian, kelemahan mendasar ")
    add_styled_run(p4, "Gaussian filter", italic=True)
    add_styled_run(p4, " terletak pada sifatnya yang ")
    add_styled_run(p4, "isotropic", italic=True)
    add_styled_run(p4, " (serba sama ke segala arah spasial) tanpa memperhitungkan variasi intensitas fotometrik antar-piksel. Ketika ")
    add_styled_run(p4, "kernel", italic=True)
    add_styled_run(p4, " konvolusi ")
    add_styled_run(p4, "Gaussian", italic=True)
    add_styled_run(p4, " melintasi perbatasan antara bayangan siluet kardiak yang radioopak dan parenkim paru yang radiolusen, nilai intensitas dari kedua kompartemen jaringan tersebut akan diratakan secara merata. Fenomena ini memicu efek pengaburan tepi (")
    add_styled_run(p4, "edge blurring", italic=True)
    add_styled_run(p4, ") yang parah, melandaikan gradien transisi tepi kardioperikardial (")
    add_styled_run(p4, "cardiopericardial border", italic=True)
    add_styled_run(p4, "), serta mengikis ketajaman garis batas sinus kostofrenikus (")
    add_styled_run(p4, "costophrenic angle", italic=True)
    add_styled_run(p4, ") dan dinding lateral sangkar toraks. Akibatnya, ketidakpastian spasial dalam penentuan titik kontur terluar jantung (")
    add_styled_run(p4, "maximum right diameter", italic=True)
    add_styled_run(p4, " dan ")
    add_styled_run(p4, "maximum left diameter", italic=True)
    add_styled_run(p4, ") meningkat secara signifikan.")

    # 9. Paragraf 5: Peran Baseline Pembanding
    p5 = create_p()
    format_p(p5)
    add_styled_run(p5, "Dalam penelitian ini, ")
    add_styled_run(p5, "Gaussian filter", italic=True)
    add_styled_run(p5, " diikutsertakan sebagai metode pembanding dasar (")
    add_styled_run(p5, "classical baseline", italic=True)
    add_styled_run(p5, ") tanpa kemampuan preservasi tepi (")
    add_styled_run(p5, "non-edge-preserving", italic=True)
    add_styled_run(p5, "). Kehadiran ")
    add_styled_run(p5, "Gaussian filter", italic=True)
    add_styled_run(p5, " berfungsi untuk membuktikan secara empiris hipotesis penelitian bahwa reduksi derau konvensional yang semata-mata mengandalkan kehalusan spasial justru merusak integritas batas anatomis kritis dan memperbesar galat pengukuran ")
    add_styled_run(p5, "Cardiothoracic Ratio", italic=True)
    add_styled_run(p5, " (CTR), berbeda dengan metode penapisan adaptif berbasis pelestarian tepi (")
    add_styled_run(p5, "Bilateral Filter", italic=True)
    add_styled_run(p5, ") maupun restorasi berbasis pembelajaran mendalam (")
    add_styled_run(p5, "deep learning", italic=True)
    add_styled_run(p5, " DnCNN).")

doc_path = r'd:\laragon\www\belajar-laravel\DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx'
test_doc_path = r'd:\laragon\www\belajar-laravel\test_241_result.docx'

doc = docx.Document(doc_path)
start_idx = None
end_idx = None

for i, p in enumerate(doc.paragraphs):
    if '2.4.1' in p.text:
        start_idx = i
    elif '2.4.2' in p.text:
        end_idx = i
        break

print(f"Original 2.4.1 range: p[{start_idx}] to p[{end_idx}]")
target_before = doc.paragraphs[end_idx]

# Insert new paragraphs
build_section_241(target_before)

# Remove old paragraphs (start_idx to end_idx - 1)
old_p_elms = [doc.paragraphs[k]._p for k in range(start_idx, end_idx)]
for elm in old_p_elms:
    elm.getparent().remove(elm)

doc.save(test_doc_path)
print(f"Saved test document to {test_doc_path}")
