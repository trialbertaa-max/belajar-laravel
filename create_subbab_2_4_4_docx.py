import os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_run(p, text, bold=False, italic=False, size_pt=12):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt)
    return r

def main():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.18)    # 3 cm
        s.bottom_margin = Inches(1.18) # 3 cm
        s.left_margin = Inches(1.57)   # 4 cm
        s.right_margin = Inches(1.18)  # 3 cm

    # Heading 2.4.4
    h = doc.add_paragraph()
    format_p(h, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_run(h, "2.4.4\tDomain Shift dan Tantangan Deep Denoising pada Radiografi Medis", bold=True)

    # Paragraf 1
    p1 = doc.add_paragraph()
    format_p(p1)
    add_run(p1, "Meskipun arsitektur DnCNN memiliki kapabilitas restorasi yang sangat tinggi, penerapannya secara langsung pada citra radiografi toraks menghadapi tantangan mendasar yang dikenal sebagai pergeseran domain (")
    add_run(p1, "domain shift", italic=True)
    add_run(p1, ") atau ketidakcocokan domain (")
    add_run(p1, "domain mismatch", italic=True)
    add_run(p1, ") (Guan & Liu, 2022). Model DnCNN orisinal dikembangkan dan dilatih menggunakan dataset citra alami fotografis (seperti dataset BSD500) dengan asumsi derau sintetis aditif bebas sinyal (")
    add_run(p1, "Additive White Gaussian Noise", italic=True)
    add_run(p1, " atau AWGN). Karakteristik statistik citra alami tersebut memiliki perbedaan fisik yang substansial dibandingkan dengan radiografi toraks medis, di mana derau radiasi sinar-X didominasi oleh fluktuasi foton Poisson yang bergantung sinyal (")
    add_run(p1, "signal-dependent quantum mottle", italic=True)
    add_run(p1, ") dan citra klinis memiliki rentang kedalaman fotometrik tinggi (12-bit atau 16-bit) dengan tumpang tindih struktur anatomi proyeksi yang sangat kompleks.")

    # Paragraf 2
    p2 = doc.add_paragraph()
    format_p(p2)
    add_run(p2, "Untuk mengatasi kesenjangan domain tersebut, penelitian ini menerapkan pendekatan adaptasi domain melalui strategi ")
    add_run(p2, "transfer learning", italic=True)
    add_run(p2, " dan penyesuaian bobot (")
    add_run(p2, "fine-tuning", italic=True)
    add_run(p2, "). Bobot awal model dasar (")
    add_run(p2, "pretrained weights", italic=True)
    add_run(p2, ") dieksploitasi untuk mendeteksi fitur spasial primitif (seperti detektor tepi dan tekstur lokal), yang kemudian dilatih ulang secara terarah menggunakan dataset radiografi toraks digital (seperti JSRT dan VinDr-CXR) yang disuntikkan model derau campuran Poisson-Gaussian terkontrol.")

    # Paragraf 3: Pengantar Risiko
    p3 = doc.add_paragraph()
    format_p(p3)
    add_run(p3, "Dalam perancangan dan evaluasi restorasi berbasis deep learning untuk citra medis, terdapat tiga risiko metodologis kritis yang wajib diantisipasi dan dimitigasi:")

    # Butir a, b, c
    risks = [
        ("a.", "Pencegahan Kebocoran Data (Data Leakage Prevention): Pembagian dataset pelatihan, validasi, dan pengujian wajib dilakukan secara ketat berbasis identitas pasien (patient-wise split), bukan pemisahan acak berbasis citra atau patch. Hal ini krusial guna memastikan model tidak mengevaluasi citra dari pasien yang anatomisnya telah dipelajari pada fase pelatihan, yang dapat menghasilkan evaluasi performa yang overestimasi secara semu (optimistic bias)."),
        ("b.", "Risiko Penghalusan Berlebih (Over-Smoothing Risk): Konvolusi dalam berisiko menginterpretasikan struktur mikro anatomi halus berkontras rendah (seperti corakan vaskular perifer paru atau kontur tipis kardioperikardial) sebagai komponen derau, sehingga struktur tersebut terhapus secara berlebihan. Pada konteks pengukuran CTR, penghalusan berlebih pada apeks kardiak lateral dapat mengikis batas ventrikel kiri, menyebabkan diameter jantung terukur menyempit secara artifisial."),
        ("c.", "Risiko Halusinasi Rekonstruksi (Hallucination Risk): Sebagai jaringan saraf non-linear dalam, model memiliki risiko memunculkan pola sintetis palsu (hallucination artifacts) yang tidak bersumber dari transmisi radiasi fisik asli (Bhadra et al., 2021). Pada radiografi toraks, halusinasi dapat mengubah kurvatur batas lateral kardiak atau membentuk garis semu yang menyerupai infiltrat paru, yang berpotensi fatal bagi akurasi pengukuran diagnostik biometri. Oleh karena itu, pengujian performa DnCNN dalam penelitian ini tidak semata dievaluasi melalui metrik kualitas umum (PSNR dan SSIM), melainkan wajib divalidasi langsung melalui metrik berbasis tugas (task-based evaluation) terhadap galat pengukuran CTR dan pelestarian batas kontur anatomi (boundary preservation).")
    ]

    for letter, text in risks:
        p_r = doc.add_paragraph()
        format_p(p_r, space_after_pt=1, space_before_pt=0, first_line_indent_pt=14.17)
        add_run(p_r, letter + "\t")
        words = text.split(" ")
        for w_idx, w in enumerate(words):
            sp = " " if w_idx < len(words) - 1 else ""
            clean_w = w.strip("(),:.")
            is_it = clean_w in [
                "Data", "Leakage", "Prevention", "patient-wise", "split", "patch", "optimistic", "bias",
                "Over-Smoothing", "Risk", "Hallucination", "artifacts", "task-based", "evaluation",
                "boundary", "preservation"
            ]
            add_run(p_r, w + sp, italic=is_it)

    out_file = r'd:\laragon\www\belajar-laravel\SUBBAB_2_4_4_DOMAIN_SHIFT.docx'
    doc.save(out_file)
    print("Successfully generated standalone Subbab 2.4.4 docx:", out_file)

if __name__ == '__main__':
    main()

