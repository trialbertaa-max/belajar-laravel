import docx
import shutil

doc_path = r'DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx'
test_doc_path = r'test_transitions_result.docx'

shutil.copy2(doc_path, test_doc_path)
doc = docx.Document(test_doc_path)

print(f"Total paragraphs before: {len(doc.paragraphs)}")

# 1. Gambar 2.1 (p[146])
p146 = doc.paragraphs[146]
print("Old p146:", p146.text)
for r in p146.runs:
    if 'sebagaimana diilustrasikan pada Gambar 2.1' in r.text:
        r.text = r.text.replace(', sebagaimana diilustrasikan pada Gambar 2.1.', ' (Gambar 2.1).')
print("New p146:", p146.text)

# 2. Gambar 2.2 (p[193] and p[194])
p193 = doc.paragraphs[193]
p194 = doc.paragraphs[194]
print("Old p193:", p193.text[-120:])
print("Old p194:", p194.text)

# Add transition to end of p193
# run[46] in p193 ends with ').'
for r in reversed(p193.runs):
    if ').' in r.text:
        r.text = r.text.replace(').', '), dengan visualisasi prinsip kerja penapisan bilateral yang disajikan pada Gambar 2.2.')
        break
print("New p193:", p193.text[-160:])

# Remove p194 element
p194_elm = p194._p
p194_elm.getparent().remove(p194_elm)
print("Removed standalone p194")

# Re-scan document to get accurate paragraph indices after removing p194
# Find p206 (poin 3), p207 (old standalone), and p210 (last p of 2.4.3)
p_layer3 = None
p_standalone23 = None
p_bridge244 = None

for i, p in enumerate(doc.paragraphs):
    if 'Lapisan Rekonstruksi Akhir' in p.text:
        p_layer3 = p
    elif 'Mekanisme pemisahan sisa derau dan konfigurasi hierarkis' in p.text:
        p_standalone23 = p
    elif 'Meskipun arsitektur DnCNN menunjukkan performa' in p.text:
        p_bridge244 = p

print("Found p_layer3:", p_layer3 is not None)
print("Found p_standalone23:", p_standalone23 is not None)
print("Found p_bridge244:", p_bridge244 is not None)

# 3. Gambar 2.3: update p_layer3 to reference (Gambar 2.3) and remove p_standalone23
for r in reversed(p_layer3.runs):
    if ').' in r.text:
        r.text = r.text.replace(').', ') (Gambar 2.3).')
        break
print("New p_layer3:", p_layer3.text)

p_standalone23_elm = p_standalone23._p
p_standalone23_elm.getparent().remove(p_standalone23_elm)
print("Removed standalone p_standalone23")

# 4. Kalimat penutup sebelum 2.4.4 (p_bridge244)
print("Old p_bridge244:", p_bridge244.text[-120:])
for r in reversed(p_bridge244.runs):
    if 'sebagaimana dibahas secara mendalam pada Subbab 2.4.4' in r.text:
        r.text = r.text.replace(', sebagaimana dibahas secara mendalam pada Subbab 2.4.4.', ' guna menjamin keandalan restorasi batas anatomis kardioperikardial tersebut.')
        break
print("New p_bridge244:", p_bridge244.text[-150:])

doc.save(test_doc_path)
print(f"Total paragraphs after: {len(doc.paragraphs)}")
print("Saved to", test_doc_path)
