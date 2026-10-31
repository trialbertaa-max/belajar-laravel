import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, size_pt=12):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt)
    return r

def main():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.18)    # 3 cm
        section.bottom_margin = Inches(1.18) # 3 cm
        section.left_margin = Inches(1.57)   # 4 cm
        section.right_margin = Inches(1.18)  # 3 cm

    # --- JUDUL SUBBAB 3.3 ---
    h2 = doc.add_paragraph()
    format_p(h2, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h2, "3.3\tBahan dan Alat", bold=True, size_pt=12)

    # Paragraf Pengantar
    p_intro = doc.add_paragraph()
    format_p(p_intro)
    add_styled_run(p_intro, "Peralatan yang digunakan dalam penelitian ini diklasifikasikan ke dalam tiga kelompok utama, yaitu perangkat keras, perangkat lunak ilmiah untuk komputasi eksperimen, dan perangkat lunak demonstrator stasiun kerja medis:")

    # 1. Perangkat Keras
    p_hw = doc.add_paragraph()
    format_p(p_hw, space_after_pt=2, first_line_indent_pt=0)
    add_styled_run(p_hw, "1.\tPerangkat Keras (", bold=False)
    add_styled_run(p_hw, "Hardware", italic=True)
    add_styled_run(p_hw, "):")

    hw_items = [
        ("a.", "Unit komputer pengembang (workstation): Prosesor multi-inti (Intel Core i7 / AMD Ryzen 7), RAM 16 GB, media penyimpanan SSD NVMe 512 GB/1 TB, dan unit pengolah grafis (GPU NVIDIA GeForce RTX Series berarsitektur CUDA) untuk akselerasi pelatihan model deep learning serta komputasi penapisan citra."),
        ("b.", "Perangkat penampil dan interaksi: Monitor panel IPS beresolusi minimal Full HD (1920 \u00d7 1080 piksel) dengan reproduksi kontras tinggi dan tetikus optik presisi untuk penarikan garis ukur anatomi kardioperikardial."),
        ("c.", "Modalitas radiografi digital klinis: Unit Digital Radiography (DR) tipe Fujifilm FDR Xair yang terintegrasi dengan Flat-Panel Detector (FPD) FDR D-EVO II pada Instalasi Radiologi CDC Surabaya, terhubung ke jaringan lokal melalui protokol TCP/IP.")
    ]
    for code, text in hw_items:
        p = doc.add_paragraph()
        format_p(p, space_after_pt=2, first_line_indent_pt=14.17)
        add_styled_run(p, code + "\t")
        # Check italic terms
        words = text.split(" ")
        for w_idx, w in enumerate(words):
            sp = " " if w_idx < len(words) - 1 else ""
            clean_w = w.strip("(),:.")
            is_it = clean_w in ["workstation", "deep", "learning", "Full", "HD", "Digital", "Radiography", "Flat-Panel", "Detector"]
            add_styled_run(p, w + sp, italic=is_it)

    # 2. Perangkat Lunak Ilmiah
    p_sw = doc.add_paragraph()
    format_p(p_sw, space_after_pt=2, first_line_indent_pt=0)
    add_styled_run(p_sw, "2.\tPerangkat Lunak Ilmiah (", bold=False)
    add_styled_run(p_sw, "Scientific Software", italic=True)
    add_styled_run(p_sw, "):")

    sw_items = [
        ("a.", "Bahasa pemrograman Python (versi 3.10+) sebagai lingkungan komputasi ilmiah utama."),
        ("b.", "Pustaka PyTorch untuk perancangan arsitektur, pelatihan, fine-tuning, dan inferensi model jaringan saraf residual DnCNN."),
        ("c.", "Pustaka OpenCV untuk implementasi algoritma penapisan spasial Bilateral Filter dan operasi morfologi citra."),
        ("d.", "Pustaka pydicom untuk pembacaan berkas citra radiografi medis format DICOM, ekstraksi metadata heksadesimal, dan penanganan polaritas fotometrik."),
        ("e.", "Pustaka scikit-image untuk kalkulasi metrik fidelitas citra global (PSNR dan SSIM)."),
        ("f.", "Pustaka SciPy dan statsmodels untuk analisis matriks numerik, pemodelan sebaran derau, dan uji statistik inferensial (uji normalitas, uji komparasi repeated measures, serta uji post-hoc dengan koreksi Holm).")
    ]
    for code, text in sw_items:
        p = doc.add_paragraph()
        format_p(p, space_after_pt=2, first_line_indent_pt=14.17)
        add_styled_run(p, code + "\t")
        words = text.split(" ")
        for w_idx, w in enumerate(words):
            sp = " " if w_idx < len(words) - 1 else ""
            clean_w = w.strip("(),:.")
            is_it = clean_w in ["fine-tuning", "Bilateral", "Filter", "repeated", "measures", "post-hoc"]
            add_styled_run(p, w + sp, italic=is_it)

    # 3. Perangkat Lunak Demonstrator
    p_dm = doc.add_paragraph()
    format_p(p_dm, space_after_pt=2, first_line_indent_pt=0)
    add_styled_run(p_dm, "3.\tPerangkat Lunak Demonstrator (", bold=False)
    add_styled_run(p_dm, "Demonstrator Software", italic=True)
    add_styled_run(p_dm, "):")

    dm_items = [
        ("a.", "Lingkungan peladen lokal (local server environment): Laragon Stack dengan interpreter PHP 8.2+ dan sistem basis data relasional MySQL/SQLite."),
        ("b.", "Framework Laravel 11 di sisi backend untuk manajemen perutean, pengelolaan berkas citra medis, dan integrasi pemanggilan API restorasi."),
        ("c.", "Pustaka pynetdicom sebagai modul layanan latar belakang (daemon) DICOM C-STORE SCP untuk menangkap aliran transmisi citra dari modalitas rontgen digital via jaringan TCP/IP."),
        ("d.", "Antarmuka penampil web (web viewer): Modul penampil berbasis lapisan grafis vektor skalabel (SVG Stage Layer) untuk rendering citra radiografi, penarikan garis ukur CTR metode Danzer 1919, serta ekspor citra beranotasi format DICOM Secondary Capture.")
    ]
    for code, text in dm_items:
        p = doc.add_paragraph()
        format_p(p, space_after_pt=2, first_line_indent_pt=14.17)
        add_styled_run(p, code + "\t")
        words = text.split(" ")
        for w_idx, w in enumerate(words):
            sp = " " if w_idx < len(words) - 1 else ""
            clean_w = w.strip("(),:.")
            is_it = clean_w in ["local", "server", "environment", "Framework", "backend", "daemon", "web", "viewer", "rendering", "Secondary", "Capture"]
            add_styled_run(p, w + sp, italic=is_it)

    out_path = r'd:\laragon\www\belajar-laravel\SUBBAB_3_3_BAHAN_DAN_ALAT.docx'
    doc.save(out_path)
    print("Successfully generated:", out_path)

if __name__ == '__main__':
    main()

