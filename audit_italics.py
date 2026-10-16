import docx
import re

doc = docx.Document('DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU.docx')

# List of common foreign/technical words that should be italicized in an Indonesian thesis
foreign_terms = [
    'hardware', 'software', 'dataset', 'baseline', 'noisy', 'clean', 'real-time', 'real time',
    'deep learning', 'machine learning', 'framework', 'ground truth', 'feature maps', 'feature map',
    'batch normalization', 'residual learning', 'convolutional neural network', 'neural network',
    'web-based', 'web based', 'workstation', 'zero-footprint', 'zero footprint', 'stage layer',
    'viewer', 'display', 'backend', 'frontend', 'daemon', 'service', 'listener', 'handler',
    'inverse affine mapping', 'lockstep', 'zoom-in', 'zoom-out', 'pan', 'drag', 'flip',
    'click-to-select', 'split-curtain', 'slider', 'difference map', 'secondary capture',
    'burn-in', 'overlay', 'accession number', 'study instance uid', 'sop class uid',
    'transfer syntax', 'middleware', 'stack', 'open source', 'server', 'desktop', 'preprocessing',
    'denoising', 'edge-preserving', 'edge preserving', 'spatial filter', 'spatial domain',
    'range domain', 'patch', 'self-similarity', 'kernel', 'convolution', 'residual', 'pooling',
    'stride', 'padding', 'loss function', 'mean squared error', 'peak signal-to-noise ratio',
    'structural similarity', 'limits of agreement', 'posteroanterior', 'anteroposterior',
    'imager pixel spacing', 'pixel spacing', 'window width', 'window level', 'rescale slope',
    'rescale intercept', 'lookup table', 'lut', 'photometric interpretation', 'monochrome',
    'midline', 'true positive', 'false positive', 'bias', 'grid', 'array', 'offline', 'online'
]

findings = []

for p_idx, p in enumerate(doc.paragraphs):
    if p_idx < 104:  # Skip front matter (cover, daftar isi, dll)
        continue
    if p_idx >= 310: # Skip daftar pustaka
        continue
    
    text = p.text
    # Check each run in paragraph
    for r in p.runs:
        r_text = r.text
        if not r_text.strip():
            continue
        
        # If run is NOT italic
        if not r.italic:
            # Check if any foreign term is in this non-italic run
            for term in foreign_terms:
                # Use regex word boundary
                pattern = r'\b' + re.escape(term) + r'\b'
                matches = re.findall(pattern, r_text, re.IGNORECASE)
                if matches:
                    findings.append({
                        'p_idx': p_idx,
                        'term': matches[0],
                        'snippet': p.text[:100],
                        'style': p.style.name
                    })

print(f"Total non-italic foreign term occurrences found: {len(findings)}")
# Group findings by paragraph
from collections import defaultdict
p_grouped = defaultdict(list)
for f in findings:
    p_grouped[f['p_idx']].append(f['term'])

for p_idx, terms in sorted(p_grouped.items()):
    p_text = doc.paragraphs[p_idx].text.strip()
    unique_terms = list(dict.fromkeys(terms))
    print(f"\nP[{p_idx}] ({doc.paragraphs[p_idx].style.name}):")
    print(f"  Terms: {', '.join(unique_terms)}")
    print(f"  Text: {p_text[:120]}...")
