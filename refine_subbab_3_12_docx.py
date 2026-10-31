import os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False, subscript=False, superscript=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic
    r.font.subscript = subscript
    r.font.superscript = superscript

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_run(p, text, bold=False, italic=False, size_pt=12, subscript=False, superscript=False):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt, subscript=subscript, superscript=superscript)
    return r

def main():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.18)    # 3 cm
        s.bottom_margin = Inches(1.18) # 3 cm
        s.left_margin = Inches(1.57)   # 4 cm
        s.right_margin = Inches(1.18)  # 3 cm

    # Heading 3.12
    h = doc.add_paragraph()
    format_p(h, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_run(h, '3.12\tAnalisis Statistik Inferensial', bold=True)

    # Paragraf Pengantar
    p1 = doc.add_paragraph()
    format_p(p1)
    add_run(p1, 'Uji statistik inferensial dirancang untuk menguji signifikansi perbedaan performa antar-metode restorasi citra. Variabel keluaran penelitian dikelompokkan menjadi luaran primer (')
    add_run(p1, 'primary outcome', italic=True)
    add_run(p1, ') berupa galat biometri (MAE CTR) serta luaran sekunder (')
    add_run(p1, 'secondary outcomes', italic=True)
    add_run(p1, ') yang mencakup metrik kualitas citra (PSNR dan SSIM), galat batas anatomi (DSC, HD95, MSD), dan waktu komputasi (')
    add_run(p1, 'runtime', italic=True)
    add_run(p1, '). Prosedur analisis statistik dieksekusi melalui tahapan berikut:')

    # Poin 1: Normalitas
    p_pt1 = doc.add_paragraph()
    format_p(p_pt1, space_after_pt=1, space_before_pt=0, first_line_indent_pt=14.17)
    add_run(p_pt1, '1.\tUji Normalitas: Distribusi data galat pengukuran diuji menggunakan uji ')
    add_run(p_pt1, 'Shapiro-Wilk', italic=True)
    add_run(p_pt1, ' pada tingkat signifikansi ')
    add_run(p_pt1, '\u03b1', italic=True)
    add_run(p_pt1, ' = 0,05.')

    # Poin 2: RM-ANOVA / Friedman
    p_pt2 = doc.add_paragraph()
    format_p(p_pt2, space_after_pt=1, space_before_pt=0, first_line_indent_pt=14.17)
    add_run(p_pt2, '2.\tUji Komparasi Multikelompok Berpasangan: Karena setiap perlakuan dievaluasi pada subjek citra yang sama (')
    add_run(p_pt2, 'within-subject design', italic=True)
    add_run(p_pt2, '), komparasi antar-kondisi (')
    add_run(p_pt2, 'clean reference', italic=True)
    add_run(p_pt2, ', ')
    add_run(p_pt2, 'noisy baseline', italic=True)
    add_run(p_pt2, ', Gaussian, Bilateral Filter, dan DnCNN) dianalisis menggunakan ')
    add_run(p_pt2, 'Repeated Measures ANOVA', italic=True)
    add_run(p_pt2, ' jika data berdistribusi normal, atau uji non-parametrik ')
    add_run(p_pt2, 'Friedman', italic=True)
    add_run(p_pt2, ' jika asumsi normalitas tidak terpenuhi. Besaran efek (')
    add_run(p_pt2, 'effect size', italic=True)
    add_run(p_pt2, ') dilaporkan menggunakan ')
    add_run(p_pt2, 'partial eta-squared', italic=True)
    add_run(p_pt2, ' (')
    add_run(p_pt2, '\u03b7', italic=True)
    add_run(p_pt2, 'p', italic=True, subscript=True)
    add_run(p_pt2, '\u00b2', superscript=True)
    add_run(p_pt2, ') atau ')
    add_run(p_pt2, "Kendall's W", italic=True)
    add_run(p_pt2, ' pada rentang interval kepercayaan 95% (95% CI).')

    # Poin 3: Post-Hoc Holm
    p_pt3 = doc.add_paragraph()
    format_p(p_pt3, space_after_pt=1, space_before_pt=0, first_line_indent_pt=14.17)
    add_run(p_pt3, '3.\tUji Lanjut Berpasangan (')
    add_run(p_pt3, 'Post-Hoc', italic=True)
    add_run(p_pt3, '): Apabila uji komparasi multikelompok menunjukkan perbedaan yang signifikan (')
    add_run(p_pt3, 'p', italic=True)
    add_run(p_pt3, ' < 0,05), analisis perbandingan berpasangan dilanjutkan menggunakan ')
    add_run(p_pt3, 'paired t-test', italic=True)
    add_run(p_pt3, ' (untuk data normal) atau ')
    add_run(p_pt3, 'Wilcoxon signed-rank test', italic=True)
    add_run(p_pt3, ' (untuk data non-parametrik) dengan penyesuaian nilai ')
    add_run(p_pt3, 'p', italic=True)
    add_run(p_pt3, ' menggunakan koreksi Holm (Holm, 1979) guna mengendalikan tingkat kesalahan tipe I (')
    add_run(p_pt3, 'family-wise error rate', italic=True)
    add_run(p_pt3, ').')

    # Poin 4: Bland-Altman & Pearson/Spearman
    p_pt4 = doc.add_paragraph()
    format_p(p_pt4, space_after_pt=2, space_before_pt=0, first_line_indent_pt=14.17)
    add_run(p_pt4, '4.\tAnalisis Kesepakatan Klinis (')
    add_run(p_pt4, 'Agreement Analysis', italic=True)
    add_run(p_pt4, '): Evaluasi kesepakatan pengukuran biometri CTR antar-metode terhadap standar referensi dinilai menggunakan metode ')
    add_run(p_pt4, 'Bland-Altman', italic=True)
    add_run(p_pt4, ' untuk menentukan deviasi rata-rata (')
    add_run(p_pt4, 'Bias', italic=True)
    add_run(p_pt4, ') dan ')
    add_run(p_pt4, 'Limits of Agreement', italic=True)
    add_run(p_pt4, ' (LoA). Analisis korelasi Pearson atau Spearman digunakan sebagai pengujian pendukung untuk menilai linearitas asosiasi.')

    out_file = r'd:\laragon\www\belajar-laravel\SUBBAB_3_12_ANALISIS_STATISTIK.docx'
    doc.save(out_file)
    print('Successfully refined Subbab 3.12:', out_file)

if __name__ == '__main__':
    main()

