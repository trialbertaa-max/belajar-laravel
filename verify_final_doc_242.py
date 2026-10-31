import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc_path = r'd:\laragon\www\belajar-laravel\DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx'
doc = docx.Document(doc_path)

print(f"Total paragraphs in document: {len(doc.paragraphs)}")

# Check Gambar 2.2 and Gambar 2.3
fig2_2_found = False
fig2_3_found = False

for i, p in enumerate(doc.paragraphs):
    if 'Gambar 2. 2' in p.text:
        prev_p = doc.paragraphs[i-1]
        drawings = prev_p._p.xpath('.//w:drawing')
        print(f"Gambar 2.2 at p[{i}]: Caption='{p.text}', Drawing in prev_p: {len(drawings) > 0}")
        fig2_2_found = True
    elif 'Gambar 2. 3' in p.text:
        prev_p = doc.paragraphs[i-1]
        drawings = prev_p._p.xpath('.//w:drawing')
        print(f"Gambar 2.3 at p[{i}]: Caption='{p.text}', Drawing in prev_p: {len(drawings) > 0}")
        fig2_3_found = True

# Check 2.4.1, 2.4.2, 2.4.3 headers
p_241 = None
p_242 = None
p_243 = None

for i, p in enumerate(doc.paragraphs):
    if '2.4.1' in p.text: p_241 = (i, p.text)
    elif '2.4.2' in p.text: p_242 = (i, p.text)
    elif '2.4.3' in p.text: p_243 = (i, p.text)

print(f"2.4.1 header: p[{p_241[0]}] -> '{p_241[1]}'")
print(f"2.4.2 header: p[{p_242[0]}] -> '{p_242[1]}'")
print(f"2.4.3 header: p[{p_243[0]}] -> '{p_243[1]}'")

# Inspect all paragraphs between 2.4.2 and 2.4.3
print(f"\n--- Paragraphs in 2.4.2 (p[{p_242[0]}] to p[{p_243[0]}]) ---")
for idx in range(p_242[0], p_243[0]):
    p = doc.paragraphs[idx]
    omaths = p._p.xpath('.//*[local-name()="oMath"]')
    drawings = p._p.xpath('.//w:drawing')
    print(f"P{idx:03d} (style={p.style.name}, math={len(omaths)}, draw={len(drawings)}): {p.text[:70]}")
