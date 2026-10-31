import os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import lxml.etree as ET

xslt_path = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
xslt = ET.parse(xslt_path)
transform = ET.XSLT(xslt)

def mml_to_omml_element(mml_str):
    mml_tree = ET.fromstring(mml_str)
    omml_tree = transform(mml_tree)
    root = omml_tree.getroot()
    if root.tag.endswith('oMath'):
        para = ET.Element('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara')
        para_pr = ET.SubElement(para, '{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathParaPr')
        jc = ET.SubElement(para_pr, '{http://schemas.openxmlformats.org/officeDocument/2006/math}jc')
        jc.set('{http://schemas.openxmlformats.org/officeDocument/2006/math}val', 'center')
        para.append(root)
        return para
    return root

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False, subscript=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic
    r.font.subscript = subscript

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, size_pt=12, subscript=False):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt, subscript=subscript)
    return r

def add_subscript_symbol(p, base, sub):
    add_styled_run(p, base, italic=True)
    add_styled_run(p, sub, italic=True, subscript=True)

def main():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.18)    # 3 cm
        s.bottom_margin = Inches(1.18) # 3 cm
        s.left_margin = Inches(1.57)   # 4 cm
        s.right_margin = Inches(1.18)  # 3 cm

    # Heading 3.8
    h = doc.add_paragraph()
    format_p(h, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h, '3.8\tPelatihan DnCNN dan Optimasi Bilateral Filter', bold=True)

    # Paragraf 1: Kondisi Perlakuan
    p1 = doc.add_paragraph()
    format_p(p1)
    add_styled_run(p1, 'Eksperimen restorasi citra dievaluasi pada lima kondisi perlakuan komparatif, yaitu citra acuan (')
    add_styled_run(p1, 'clean reference', italic=True)
    add_styled_run(p1, '), citra berderau tanpa restorasi (')
    add_styled_run(p1, 'noisy baseline', italic=True)
    add_styled_run(p1, '), citra terestorasi Gaussian filter sebagai ')
    add_styled_run(p1, 'baseline', italic=True)
    add_styled_run(p1, ' linear konvensional, citra terestorasi Bilateral Filter, dan citra terestorasi DnCNN.')

    # Paragraf 2: Optimasi Bilateral Filter (dengan sigma_s dan sigma_r bersubskrip rapi)
    p2 = doc.add_paragraph()
    format_p(p2)
    add_styled_run(p2, 'Untuk Bilateral Filter, penentuan parameter spasial (')
    add_subscript_symbol(p2, '\u03c3', 's')
    add_styled_run(p2, ') dan rentang intensitas (')
    add_subscript_symbol(p2, '\u03c3', 'r')
    add_styled_run(p2, ') dilakukan secara sistematis menggunakan pencarian kisi (')
    add_styled_run(p2, 'grid search', italic=True)
    add_styled_run(p2, ') pada ')
    add_styled_run(p2, 'validation set', italic=True)
    add_styled_run(p2, ' (bukan angka acak) untuk menemukan kombinasi parameter yang menghasilkan nilai rata-rata PSNR dan SSIM tertinggi sebelum diterapkan pada pengujian akurasi biometri CTR.')

    # Paragraf 3: Model DnCNN & Training Setup
    p3 = doc.add_paragraph()
    format_p(p3)
    add_styled_run(p3, 'Untuk model DnCNN, arsitektur yang digunakan tersusun atas 17 lapisan konvolusi berbasis ')
    add_styled_run(p3, 'residual learning', italic=True)
    add_styled_run(p3, ' dengan kombinasi ')
    add_styled_run(p3, 'Batch Normalization', italic=True)
    add_styled_run(p3, ' (BN) dan ')
    add_styled_run(p3, 'Rectified Linear Unit', italic=True)
    add_styled_run(p3, ' (ReLU) (Zhang et al., 2017). Dataset radiografi toraks PA dibagi dengan proporsi 70% data latih (')
    add_styled_run(p3, 'training set', italic=True)
    add_styled_run(p3, '), 15% data validasi (')
    add_styled_run(p3, 'validation set', italic=True)
    add_styled_run(p3, '), dan 15% data uji (')
    add_styled_run(p3, 'test set', italic=True)
    add_styled_run(p3, ') secara berbasis pasien (')
    add_styled_run(p3, 'patient-wise split', italic=True)
    add_styled_run(p3, ') guna mencegah kebocoran data (')
    add_styled_run(p3, 'data leakage', italic=True)
    add_styled_run(p3, '). Citra resolusi tinggi dipotong menjadi sub-citra (')
    add_styled_run(p3, 'patches', italic=True)
    add_styled_run(p3, ') berukuran 40 \u00d7 40 piksel dengan pergeseran (')
    add_styled_run(p3, 'stride', italic=True)
    add_styled_run(p3, ') 20 piksel untuk efisiensi memori GPU dan pencegahan ')
    add_styled_run(p3, 'overfitting', italic=True)
    add_styled_run(p3, '. Pelatihan dieksekusi menggunakan ')
    add_styled_run(p3, 'optimizer', italic=True)
    add_styled_run(p3, ' Adam dengan kecepatan pembelajaran awal (')
    add_styled_run(p3, 'initial learning rate', italic=True)
    add_styled_run(p3, ') sebesar 10\u207b\u00b3 serta mengoptimalkan fungsi kerugian ')
    add_styled_run(p3, 'Mean Squared Error', italic=True)
    add_styled_run(p3, ' (MSE) atau L2 ')
    add_styled_run(p3, 'loss', italic=True)
    add_styled_run(p3, ' antara estimasi derau residual dan derau aktual sebagai berikut:')

    # Persamaan Loss Function
    mml_loss = '''<math xmlns="http://www.w3.org/1998/Math/MathML">
      <mrow>
        <mi>&#x2112;</mi><mo>(</mo><mi>&#920;</mi><mo>)</mo>
        <mo>=</mo>
        <mfrac>
          <mn>1</mn>
          <mrow><mn>2</mn><mi>N</mi></mrow>
        </mfrac>
        <munderover>
          <mo>&#8721;</mo>
          <mrow><mi>i</mi><mo>=</mo><mn>1</mn></mrow>
          <mi>N</mi>
        </munderover>
        <msubsup>
          <mrow>
            <mo>&#8214;</mo>
            <mi>&#x211B;</mi><mo>(</mo><msub><mi>y</mi><mi>i</mi></msub><mo>;</mo><mi>&#920;</mi><mo>)</mo>
            <mo>-</mo>
            <mo>(</mo><msub><mi>y</mi><mi>i</mi></msub><mo>-</mo><msub><mi>x</mi><mi>i</mi></msub><mo>)</mo>
            <mo>&#8214;</mo>
          </mrow>
          <mi>F</mi>
          <mn>2</mn>
        </msubsup>
      </mrow>
    </math>'''

    p_eq = doc.add_paragraph()
    format_p(p_eq, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, space_before_pt=4, line_spacing=1.0, first_line_indent_pt=0)
    omml_loss = mml_to_omml_element(mml_loss)
    p_eq._p.append(omml_loss)

    # Paragraf 4: Keterangan parameter & Fine-tuning domain shift
    p4 = doc.add_paragraph()
    format_p(p4)
    add_styled_run(p4, 'dengan ')
    add_styled_run(p4, '\u211B(yi; \u0398)', italic=False)
    add_styled_run(p4, ' merepresentasikan pemetaan residual (')
    add_styled_run(p4, 'residual mapping', italic=True)
    add_styled_run(p4, ') jaringan berbobot ')
    add_styled_run(p4, '\u0398', italic=True)
    add_styled_run(p4, ', ')
    add_styled_run(p4, 'yi', italic=True)
    add_styled_run(p4, ' adalah sampel ')
    add_styled_run(p4, 'patch', italic=True)
    add_styled_run(p4, ' berderau, ')
    add_styled_run(p4, 'xi', italic=True)
    add_styled_run(p4, ' adalah sampel acuan bebas derau, dan N adalah jumlah total ')
    add_styled_run(p4, 'patch', italic=True)
    add_styled_run(p4, ' latih. Model kemudian di-')
    add_styled_run(p4, 'fine-tune', italic=True)
    add_styled_run(p4, ' pada citra radiografi toraks PA terdegradasi untuk memitigasi ketidakcocokan domain (')
    add_styled_run(p4, 'domain mismatch', italic=True)
    add_styled_run(p4, ') dan risiko distorsi batas anatomi (Nazir et al., 2024).')

    out_file = r'd:\laragon\www\belajar-laravel\SUBBAB_3_8_PELATIHAN_DAN_OPTIMASI_FINAL.docx'
    try:
        doc.save(out_file)
        print('Successfully saved to:', out_file)
    except PermissionError:
        alt_file = r'd:\laragon\www\belajar-laravel\SUBBAB_3_8_PELATIHAN_DAN_OPTIMASI_FINAL_SUB.docx'
        doc.save(alt_file)
        print('Saved to alternate:', alt_file)

if __name__ == '__main__':
    main()
