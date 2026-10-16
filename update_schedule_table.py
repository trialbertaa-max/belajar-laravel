import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc_path = 'd:/laragon/www/belajar-laravel/DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU.docx'
doc = docx.Document(doc_path)

# 1. Pastikan P[222] (Subbab 3.1) sinkron 5 bulan (September 2026 - Januari 2027)
for i, p in enumerate(doc.paragraphs):
    if 'Waktu penelitian direncanakan berlangsung selama' in p.text:
        print(f"Found Waktu Penelitian at P[{i}]: {p.text}")
        if '4 bulan' in p.text:
            p.text = p.text.replace('4 bulan, dimulai pada bulan Oktober 2026', '5 bulan, dimulai pada bulan September 2026')
            print("Updated P to 5 bulan September 2026.")

# 2. Pastikan pengantar P[291] sinkron 5 bulan
for i, p in enumerate(doc.paragraphs):
    if 'Rencana pelaksanaan penelitian dirancang selama' in p.text and 'sebagaimana dijabarkan' in p.text:
        print(f"Found Pengantar Jadwal at P[{i}]: {p.text}")
        p.text = "Rencana pelaksanaan penelitian dirancang selama 5 bulan (September 2026 – Januari 2027) sebagaimana dijabarkan pada tabel berikut:"
        print(f"Updated Pengantar to: {p.text}")

# 3. Temukan Table 3 (Tabel Jadwal Penelitian)
target_tbl_idx = None
for idx, tbl in enumerate(doc.tables):
    first_row = [c.text.strip().replace('\n', ' ') for c in tbl.rows[0].cells]
    if any('kegiatan' in x.lower() for x in first_row) and any('bulan' in x.lower() or 'sep' in x.lower() for x in first_row):
        target_tbl_idx = idx
        break

print(f"Target schedule table index: {target_tbl_idx}")
old_table = doc.tables[target_tbl_idx]

# Data Kegiatan 7 Baris x 5 Bulan (September - Januari)
headers = ['No', 'Rincian Kegiatan', 'Bulan ke-1\n(Sep)', 'Bulan ke-2\n(Okt)', 'Bulan ke-3\n(Nov)', 'Bulan ke-4\n(Des)', 'Bulan ke-5\n(Jan)']

rows_data = [
    {
        'no': '1',
        'kegiatan': 'Studi literatur, kurasi dataset DICOM CDC Surabaya, & perancangan sistem',
        'active_months': [1] # Sep
    },
    {
        'no': '2',
        'kegiatan': 'Pembuatan Modul 1 (DICOM Listener SCP) & Modul 2 (Backend Laravel & Database)',
        'active_months': [1, 2] # Sep, Okt
    },
    {
        'no': '3',
        'kegiatan': 'Pengembangan Modul 3 (Engine Restorasi AI: DnCNN, NLM, Bilateral) & uji derau',
        'active_months': [2] # Okt
    },
    {
        'no': '4',
        'kegiatan': 'Pengembangan Modul 4 (Web Workstation Hyu PACS Viewer & Kaliper CTR Danzer)',
        'active_months': [2, 3] # Okt, Nov
    },
    {
        'no': '5',
        'kegiatan': 'Pengembangan Modul 5 (Pelaporan Medis & Ekspor Secondary Capture PS 3.3)',
        'active_months': [3] # Nov
    },
    {
        'no': '6',
        'kegiatan': 'Analisis akhir & evaluasi kuantitatif (metrik PSNR, SSIM, CTR, & statistik)',
        'active_months': [4] # Des
    },
    {
        'no': '7',
        'kegiatan': 'Penyusunan laporan skripsi & penyelesaian akhir naskah',
        'active_months': [4, 5] # Des, Jan
    }
]

# Buat tabel baru 8 baris x 7 kolom
new_table = doc.add_table(rows=8, cols=7)
new_table.alignment = WD_TABLE_ALIGNMENT.CENTER
new_table.style = 'Table Grid'

# Pindahkan posisi tabel baru persis ke posisi old_table
old_table._element.addprevious(new_table._element)
old_table._element.getparent().remove(old_table._element)

# Atur lebar kolom (Total ~ 14.5 cm)
col_widths = [Cm(1.0), Cm(8.0), Cm(1.1), Cm(1.1), Cm(1.1), Cm(1.1), Cm(1.1)]

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shd if any
    for existing in tcPr.findall(qn('w:shd')):
        tcPr.remove(existing)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

# Format Header (Row 0)
for c_idx, title in enumerate(headers):
    cell = new_table.cell(0, c_idx)
    cell.width = col_widths[c_idx]
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    run = p.add_run(title)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10.5)
    run.font.bold = True
    
    # Header shading abu-abu muda elegan
    set_cell_shading(cell, "F1F5F9")

# Format Rows 1 to 7
for r_idx, item in enumerate(rows_data, start=1):
    # Col 0: No
    cell_no = new_table.cell(r_idx, 0)
    cell_no.width = col_widths[0]
    cell_no.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell_no)
    p_no = cell_no.paragraphs[0]
    p_no.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_no.paragraph_format.line_spacing = 1.15
    p_no.paragraph_format.space_before = Pt(0)
    p_no.paragraph_format.space_after = Pt(0)
    r_no = p_no.add_run(item['no'])
    r_no.font.name = 'Times New Roman'
    r_no.font.size = Pt(11)
    
    # Col 1: Rincian Kegiatan
    cell_act = new_table.cell(r_idx, 1)
    cell_act.width = col_widths[1]
    cell_act.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell_act, top=100, bottom=100, left=150, right=150)
    p_act = cell_act.paragraphs[0]
    p_act.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_act.paragraph_format.line_spacing = 1.15
    p_act.paragraph_format.space_before = Pt(0)
    p_act.paragraph_format.space_after = Pt(0)
    r_act = p_act.add_run(item['kegiatan'])
    r_act.font.name = 'Times New Roman'
    r_act.font.size = Pt(11)
    
    # Col 2-6: Bulan 1-5
    for m_idx in range(1, 6):
        cell_m = new_table.cell(r_idx, m_idx + 1)
        cell_m.width = col_widths[m_idx + 1]
        cell_m.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell_m)
        p_m = cell_m.paragraphs[0]
        p_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_m.paragraph_format.line_spacing = 1.0
        p_m.paragraph_format.space_before = Pt(0)
        p_m.paragraph_format.space_after = Pt(0)
        
        if m_idx in item['active_months']:
            set_cell_shading(cell_m, "000000") # Black block standard Gantt
        else:
            set_cell_shading(cell_m, "FFFFFF")

# Simpan langsung ke dokumen
doc.save(doc_path)
print("Successfully updated Table 3.3 and proposal paragraphs in DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU.docx!")
