import os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False, subscript=False, superscript=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic
    r.font.subscript = subscript
    r.font.superscript = superscript

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, size_pt=12, subscript=False, superscript=False):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt, subscript=subscript, superscript=superscript)
    return r

def main():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.18)    # 3 cm
        s.bottom_margin = Inches(1.18) # 3 cm
        s.left_margin = Inches(1.57)   # 4 cm
        s.right_margin = Inches(1.18)  # 3 cm

    # Heading 3.13
    h = doc.add_paragraph()
    format_p(h, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h, '3.13\tPerancangan Demonstrator Web-Based DICOM Workstation', bold=True)

    # Paragraf 1: Pengantar & Batasan Masalah (BM 6)
    p1 = doc.add_paragraph()
    format_p(p1)
    add_styled_run(p1, 'Sebagai luaran implementasi praktis dan pembuktian konsep (')
    add_styled_run(p1, 'proof-of-concept', italic=True)
    add_styled_run(p1, '), sistem perangkat lunak dirancang sebagai demonstrator ')
    add_styled_run(p1, 'Web-Based DICOM Workstation', italic=True)
    add_styled_run(p1, ' mandiri guna menguji aspek interoperabilitas transmisi medis, eksekusi restorasi citra, dan kemudahan penarikan garis ukur biometri CTR. Sesuai batasan masalah penelitian, sistem ini difungsikan murni sebagai demonstrator penelitian untuk menguji kelayakan komputasional dan alur kerja klinis, bukan sebagai sistem PACS klinis tervalidasi.')

    # Paragraf 2: 4 Modul Terpangkas
    p2 = doc.add_paragraph()
    format_p(p2)
    add_styled_run(p2, 'Arsitektur sistem demonstrator disederhanakan secara terfokus menjadi empat modul fungsional utama:')

    # 4 Modul Terstruktur
    modules = [
        ("1.", "Modul Ingesti Citra (DICOM Input): Menangani penerimaan berkas radiografi toraks PA melalui layanan latar belakang C-STORE SCP (port 4242, Application Entity Title: HYU_PACS_SERVER) berbasis pustaka pynetdicom maupun pengunggahan berkas .dcm secara langsung melalui antarmuka web, dilanjutkan dengan ekstraksi metadata pasien dan larik piksel mentah (raw pixel array)."),
        ("2.", "Modul Pemrosesan Citra (Image Processing API): Layanan backend berbasis Python yang mengeksekusi inferensi algoritma restorasi citra teruji (Bilateral Filter dan model DnCNN) untuk menghasilkan citra terestorasi bebas derau secara terotomasi."),
        ("3.", "Modul Penampil dan Pengukuran (CTR Measurement Viewer): Antarmuka penampil citra medis berbasis web yang dilengkapi SVG Stage Layer dengan pemetaan matriks transformasi terbalik (inverse Current Transformation Matrix atau CTM\u207b\u00b9) sehingga titik jangkar kaliper CTR Danzer tetap terkunci secara presisi (lockstep) pada piksel anatomi saat citra dimanipulasi melalui operasi windowing, zoom, pan, rotasi, dan flip."),
        ("4.", "Modul Penyimpanan dan Ekspor Hasil (Result Storage and Export): Menyimpan nilai pengukuran rasio CTR, status klasifikasi, citra hasil restorasi, dan riwayat audit ke basis data, serta menyediakan fungsi ekspor citra beranotasi ke dalam format DICOM Secondary Capture (SC) dan PNG resolusi tinggi untuk keperluan demonstrasi.")
    ]

    for num, text in modules:
        p_mod = doc.add_paragraph()
        format_p(p_mod, space_after_pt=1, space_before_pt=0, first_line_indent_pt=14.17)
        add_styled_run(p_mod, num + '\t')
        words = text.split(' ')
        for w_idx, w in enumerate(words):
            sp = ' ' if w_idx < len(words) - 1 else ''
            clean_w = w.strip('(),:.')
            is_it = clean_w in [
                'DICOM', 'Input', 'C-STORE', 'SCP', 'port', 'Application', 'Entity', 'Title',
                'pynetdicom', 'web', 'raw', 'pixel', 'array', 'Image', 'Processing', 'API',
                'backend', 'Bilateral', 'Filter', 'DnCNN', 'CTR', 'Measurement', 'Viewer',
                'SVG', 'Stage', 'Layer', 'inverse', 'Current', 'Transformation', 'Matrix', 'CTM\u207b\u00b9',
                'lockstep', 'windowing', 'zoom', 'pan', 'flip', 'Result', 'Storage', 'and', 'Export',
                'Secondary', 'Capture', 'SC', 'PNG'
            ]
            add_styled_run(p_mod, w + sp, italic=is_it)

    # Paragraf 3: Pengantar Metrik Kinerja Kuantitatif
    p3 = doc.add_paragraph()
    format_p(p3)
    add_styled_run(p3, 'Untuk menguji kelayakan operasional sistem secara kuantitatif, evaluasi kinerja perangkat lunak dinilai melalui dua parameter kinerja teknis:')

    # 2 Parameter Kinerja
    perf_metrics = [
        ("1.", "Latensi Inferensi Restorasi (Inference / Restoration Latency): Mengukur durasi komputasi yang dibutuhkan model restorasi (Bilateral Filter dan DnCNN) dalam memproses satu citra DICOM beresolusi penuh, dinyatakan dalam milidetik (ms)."),
        ("2.", "Waktu Render Penampil (Viewer Rendering Time): Mengukur kecepatan antarmuka SVG Stage Layer dalam menampilkan citra dan memperbarui matriks koordinat transformasi affine saat interaksi perbesaran (zoom) dan pergeseran (pan), dinyatakan dalam milidetik (ms).")
    ]

    for num, text in perf_metrics:
        p_pm = doc.add_paragraph()
        format_p(p_pm, space_after_pt=1, space_before_pt=0, first_line_indent_pt=14.17)
        add_styled_run(p_pm, num + '\t')
        words = text.split(' ')
        for w_idx, w in enumerate(words):
            sp = ' ' if w_idx < len(words) - 1 else ''
            clean_w = w.strip('(),:.')
            is_it = clean_w in [
                'Inference', 'Restoration', 'Latency', 'Viewer', 'Rendering', 'Time',
                'SVG', 'Stage', 'Layer', 'affine', 'zoom', 'pan'
            ]
            add_styled_run(p_pm, w + sp, italic=is_it)

    # Gambar 3.4
    img_path = r'd:\laragon\www\belajar-laravel\Gambar_3_4_Demonstrator_DICOM_4_Modul.png'
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        format_p(p_img, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, space_before_pt=8, line_spacing=1.0, first_line_indent_pt=0)
        p_img.add_run().add_picture(img_path, width=Inches(5.4))

        # Caption Gambar 3.4
        p_cap = doc.add_paragraph()
        format_p(p_cap, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=6, space_before_pt=2, line_spacing=1.15, first_line_indent_pt=0)
        add_styled_run(p_cap, 'Gambar 3.4 Diagram Alur Kerja Perangkat Lunak Demonstrator Web-Based DICOM Workstation', bold=False)

    out_file = r'd:\laragon\www\belajar-laravel\SUBBAB_3_13_DEMONSTRATOR_DICOM.docx'
    doc.save(out_file)
    print('Successfully generated Subbab 3.13:', out_file)

if __name__ == '__main__':
    main()

