import docx

doc = docx.Document('d:/laragon/www/belajar-laravel/RUMUS_SUBBAB_2_6.docx')
print("Total paragraphs:", len(doc.paragraphs))

eq_count = 0
for idx, p in enumerate(doc.paragraphs):
    omaths = p._p.xpath('.//m:oMath')
    if omaths:
        eq_count += 1
        math_text = omaths[0].xpath('string(.)').encode('ascii', 'replace').decode('ascii')
        print(f"Eq {eq_count} at p[{idx}]: {math_text}")

print(f"Total OMML equations found: {eq_count}")
