import docx

doc = docx.Document('d:/laragon/www/belajar-laravel/DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx')

with open('updated_section24_dump.txt', 'w', encoding='utf-8') as f:
    f.write("=== UPDATED SECTION 2.4 IN PROPOSAL DOCX ===\n\n")
    for i in range(165, 215):
        if i >= len(doc.paragraphs):
            break
        p = doc.paragraphs[i]
        has_draw = 'w:drawing' in p._p.xml
        has_math = 'm:oMath' in p._p.xml or 'm:oMathPara' in p._p.xml
        f.write(f"p[{i}] style={p.style.name} math={has_math} drawing={has_draw} text={p.text[:90]!r}\n")

print("Dumped updated section to updated_section24_dump.txt")
