import docx

doc = docx.Document('recovered_$R2V56TF.asd.asd')
print("Total paragraphs in recovered:", len(doc.paragraphs))

# Let's see some paragraphs in 2.4 to see what the user had edited manually
for i, p in enumerate(doc.paragraphs):
    if '2.4' in p.text:
        print(f"p[{i}]: {p.text[:100]}")
