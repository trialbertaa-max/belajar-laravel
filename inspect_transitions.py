import docx

doc = docx.Document('DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx')

with open('inspect_transitions.txt', 'w', encoding='utf-8') as f:
    def check_context(label, query):
        for i, p in enumerate(doc.paragraphs):
            if query in p.text:
                f.write(f"\n=== Found {label} at p[{i}] ===\n")
                for k in range(max(0, i-4), min(len(doc.paragraphs), i+4)):
                    drawings = p._p.xpath('.//w:drawing')
                    f.write(f"  p[{k}] (style: {doc.paragraphs[k].style.name}, align: {doc.paragraphs[k].alignment}): {doc.paragraphs[k].text}\n")
                    for r_idx, r in enumerate(doc.paragraphs[k].runs):
                        f.write(f"      run[{r_idx}] i={r.italic} b={r.bold}: {repr(r.text)}\n")

    check_context('Gambar 2.1', 'Gambar 2. 1')
    check_context('Gambar 2.2', 'Gambar 2. 2')
    check_context('Gambar 2.3', 'Gambar 2. 3')
    check_context('Subbab 2.4.4', '2.4.4')

print("Done writing inspect_transitions.txt")
