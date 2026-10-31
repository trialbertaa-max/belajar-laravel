import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, size_pt=12):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt)
    return r

def main():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.18)    # 3 cm
        section.bottom_margin = Inches(1.18) # 3 cm
        section.left_margin = Inches(1.57)   # 4 cm
        section.right_margin = Inches(1.18)  # 3 cm

    # --- JUDUL SUBBAB 3.4 ---
    h2 = doc.add_paragraph()
    format_p(h2, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h2, "3.4\tDataset dan Kriteria Inklusi", bold=True, size_pt=12)

    # Paragraf Pengantar
    p_intro = doc.add_paragraph()
    format_p(p_intro)
    add_styled_run(p_intro, "Data yang digunakan dalam penelitian ini bersumber dari kombinasi dataset repositori standar internasional terbuka (")
    add_styled_run(p_intro, "open-access benchmark datasets", italic=True)
    add_styled_run(p_intro, ") dan dataset klinis lokal teranonimisasi:")

    # 3 Dataset (Hanya 1 tingkat penomoran: 1, 2, 3)
    datasets = [
        ("1.", "Dataset A (JSRT dan Masker Segmentasi SCR): Merupakan dataset acuan standar dari Japanese Society of Radiological Technology yang dilengkapi anotasi segmentasi anatomi Segmentation in Chest Radiographs (SCR) oleh konsensus pakar radiologi (van Ginneken et al., 2006). Dataset ini terdiri atas 247 citra radiografi toraks PA beresolusi tinggi (2048 \u00d7 2048 piksel) dengan kedalaman 12-bit. Dataset A difungsikan sebagai sumber utama ekstraksi kontur batas anatomi siluet kardiak dan rongga toraks serta pembentukan nilai acuan standar penelitian (reference standard) CTR."),
        ("2.", "Dataset B (VinDr-CXR): Merupakan dataset publik berskala besar yang dikurasi dari rumah sakit rujukan utama dan dipublikasikan melalui PhysioNet (Nguyen et al., 2022). Dataset ini menyediakan 18.000 citra radiografi toraks PA dalam format asli DICOM lengkap dengan metadata klinis asli dan anotasi konsensus dari radiolog senior. Dataset B digunakan untuk pengujian domain radiografi digital riil serta pelatihan dan penyesuaian (fine-tuning) arsitektur jaringan saraf tiruan DnCNN pada citra berkedalaman bit tinggi."),
        ("3.", "Dataset C (Data Klinis CDC Surabaya): Merupakan dataset radiografi toraks PA berformat DICOM dari pasien pemeriksaan kesehatan berkala (Medical Check-Up atau MCU) pada Instalasi Radiologi Cahaya Diagnostic Centre (CDC) Surabaya. Dataset ini difungsikan secara terbatas sebagai kelompok validasi eksternal (external validation) guna menguji ketahanan model dan fungsi interoperabilitas purwarupa stasiun kerja DICOM berbasis web pada alur operasional klinis riil. Seluruh informasi identitas pribadi pasien (seperti nama dan nomor rekam medis) telah dihilangkan secara menyeluruh melalui prosedur de-identifikasi (de-identification) untuk menjamin prinsip kerahasiaan data medis.")
    ]

    for num, text in datasets:
        p = doc.add_paragraph()
        format_p(p, space_after_pt=2, first_line_indent_pt=14.17)
        add_styled_run(p, num + "\t")
        words = text.split(" ")
        for w_idx, w in enumerate(words):
            sp = " " if w_idx < len(words) - 1 else ""
            clean_w = w.strip("(),:.")
            is_it = clean_w in ["Japanese", "Society", "of", "Radiological", "Technology", "Segmentation", "in", "Chest", "Radiographs", "et", "al.", "reference", "standard", "fine-tuning", "Medical", "Check-Up", "external", "validation", "de-identification"]
            add_styled_run(p, w + sp, italic=is_it)

    # Paragraf Kriteria Inklusi & Eksklusi (Opsi 2: Naratif Mengalir Tanpa Nomor Bertingkat)
    p_kriteria = doc.add_paragraph()
    format_p(p_kriteria, space_before_pt=4)
    add_styled_run(p_kriteria, "Kriteria inklusi pada penelitian ini mencakup citra radiografi toraks proyeksi PA posisi tegak (")
    add_styled_run(p_kriteria, "upright", italic=True)
    add_styled_run(p_kriteria, ") pasien dewasa, derajat inspirasi optimal yang memperlihatkan minimal 9 hingga 10 arkus kosta posterior di atas kubah hemidiafragma, serta kontur batas lateral jantung dan sudut sinus kostofrenikus tampak tegas. Sementara itu, kriteria eksklusi meliputi citra proyeksi AP atau oblik, citra dengan rotasi tubuh yang signifikan (ditandai dengan asimetri jarak ujung medial klavikula kanan-kiri terhadap garis prosesus spinosus vertebra torakalis), serta citra dengan kelainan patologis masif seperti efusi pleura masif bilateral atau kardiomegali derajat berat yang mengaburkan batas anatomi kardioperikardial.")

    out_path = r'd:\laragon\www\belajar-laravel\SUBBAB_3_4_DATASET.docx'
    doc.save(out_path)
    print("Successfully generated clean 3.4 docx:", out_path)

if __name__ == '__main__':
    main()
