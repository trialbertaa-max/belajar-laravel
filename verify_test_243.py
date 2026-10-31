import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx

doc = docx.Document('d:/laragon/www/belajar-laravel/test_243_result.docx')

start_idx = None
end_idx = None

for i, p in enumerate(doc.paragraphs):
    if '2.4.3' in p.text:
        start_idx = i
    elif '2.4.4' in p.text:
        end_idx = i
        break

print(f"Verified 2.4.3 range in test document: p[{start_idx}] to p[{end_idx}]")

with open('verified_test_243.txt', 'w', encoding='utf-8') as f:
    for i in range(start_idx, end_idx):
        p = doc.paragraphs[i]
        f.write(f"\n========================================\n")
        f.write(f"Paragraph {i} (style={p.style.name}, align={p.alignment}):\n")
        f.write(f"Full text: {p.text}\n")
        omaths = p._p.xpath('.//*[local-name()="oMath"]')
        f.write(f"oMath count: {len(omaths)}\n")
        for om in omaths:
            tokens = [t.text for t in om.xpath('.//*[local-name()="t"]') if t.text]
            f.write(f"  Math tokens: {' '.join(tokens)}\n")
        drawings = p._p.xpath('.//w:drawing')
        f.write(f"drawings count: {len(drawings)}\n")
        f.write(f"Runs ({len(p.runs)}):\n")
        for r_idx, r in enumerate(p.runs):
            f.write(f"  [{r_idx}] text='{r.text}' | italic={r.italic} | bold={r.bold} | sub={r.font.subscript}\n")

print("Saved verified_test_243.txt")
