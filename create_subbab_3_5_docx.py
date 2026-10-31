import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, size_pt=12):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt)
    return r

def main():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.18)    # 3 cm
        section.bottom_margin = Inches(1.18) # 3 cm
        section.left_margin = Inches(1.57)   # 4 cm
        section.right_margin = Inches(1.18)  # 3 cm

    # --- JUDUL SUBBAB 3.5 ---
    h2 = doc.add_paragraph()
    format_p(h2, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h2, "3.5\tStandar Referensi Penelitian CTR", bold=True, size_pt=12)

    # Paragraf 1: Definisi & Pembentukan Reference Standard
    p1 = doc.add_paragraph()
    format_p(p1)
    add_styled_run(p1, "Evaluasi akurasi pengukuran ")
    add_styled_run(p1, "Cardiothoracic Ratio", italic=True)
    add_styled_run(p1, " (CTR) pasca-restorasi memerlukan nilai acuan pembanding yang konsisten dan terstandarisasi. Dalam penelitian ini, nilai pengukuran pada citra berkualitas optimal tanpa derau (")
    add_styled_run(p1, "clean baseline", italic=True)
    add_styled_run(p1, ") ditetapkan sebagai standar referensi penelitian (")
    add_styled_run(p1, "reference standard", italic=True)
    add_styled_run(p1, "), bukan sebagai nilai acuan mutlak. Pada dataset tolok ukur (Dataset A: JSRT dan SCR), nilai standar referensi CTR dibentuk secara objektif melalui penentuan koordinat batas lateral kardioperikardial dan dinding toraks yang diekstraksi dari masker segmentasi konsensus pakar radiologi (van Ginneken ")
    add_styled_run(p1, "et al.", italic=True)
    add_styled_run(p1, ", 2006).")

    # Paragraf 2: Prosedur pada Data Klinis & Tujuan Objektif
    p2 = doc.add_paragraph()
    format_p(p2)
    add_styled_run(p2, "Sementara itu, pada kelompok citra klinis (Dataset C: CDC Surabaya), nilai standar referensi diperoleh melalui penarikan garis ukur pada citra ")
    add_styled_run(p2, "clean baseline", italic=True)
    add_styled_run(p2, " menggunakan kaidah baku metode Danzer (1919) yang diverifikasi oleh praktisi radiologi klinis. Penetapan standar referensi penelitian ini bertujuan menyediakan tolok ukur komparasi yang objektif dan dapat direproduksi (")
    add_styled_run(p2, "reproducible", italic=True)
    add_styled_run(p2, ") guna mengkuantifikasi galat pengukuran CTR akibat degradasi derau sintetis dan pemulihannya oleh metode restorasi.")

    out_path = r'd:\laragon\www\belajar-laravel\SUBBAB_3_5_REFERENCE_STANDARD.docx'
    doc.save(out_path)
    print("Successfully generated Subbab 3.5 docx:", out_path)

if __name__ == '__main__':
    main()

