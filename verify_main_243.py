import docx

doc = docx.Document('DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx')

start_idx = None
end_idx = None
for i, p in enumerate(doc.paragraphs):
    if '2.4.3' in p.text:
        start_idx = i
    if '2.4.4' in p.text:
        end_idx = i
        break

print(f"2.4.3 is from p[{start_idx}] to p[{end_idx}]")

with open('verified_main_243.txt', 'w', encoding='utf-8') as f:
    for idx in range(start_idx, end_idx + 2):
        p = doc.paragraphs[idx]
        omaths = p._p.xpath('.//m:oMath')
        drawings = p._p.xpath('.//w:drawing')
        f.write(f"\n========================================\n")
        f.write(f"Paragraph {idx} (style: {p.style.name}, align: {p.alignment})\n")
        f.write(f"oMath count: {len(omaths)}, drawings count: {len(drawings)}\n")
        f.write(f"Text: {p.text}\n")
        for r_idx, r in enumerate(p.runs):
            f.write(f"  [{r_idx}] i={r.italic} b={r.bold} sub={r.font.subscript} sup={r.font.superscript} : {repr(r.text)}\n")

print("Saved verified_main_243.txt")
