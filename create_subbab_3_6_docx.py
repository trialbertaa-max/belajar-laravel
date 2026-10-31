import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, size_pt=12):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt)
    return r

def main():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.18)    # 3 cm
        section.bottom_margin = Inches(1.18) # 3 cm
        section.left_margin = Inches(1.57)   # 4 cm
        section.right_margin = Inches(1.18)  # 3 cm

    # --- JUDUL SUBBAB 3.6 ---
    h2 = doc.add_paragraph()
    format_p(h2, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h2, "3.6\tPrapemrosesan Citra DICOM", bold=True, size_pt=12)

    # Paragraf Pengantar Singkat
    p_intro = doc.add_paragraph()
    format_p(p_intro)
    add_styled_run(p_intro, "Data piksel mentah dan metadata DICOM diproses awal untuk menyeragamkan representasi intensitas melalui alur pipa terstandarisasi sebagai berikut:")

    # 7 Langkah Ringkas (1 kalimat per butir)
    steps = [
        ("1.", "Pembacaan larik piksel (read pixel array): Mengekstraksi matriks data piksel mentah dari berkas DICOM beserta metadata dimensi spasial dan kedalaman bita."),
        ("2.", "Pemeriksaan interpretasi fotometrik (photometric interpretation): Memeriksa polaritas citra pada tag (0028, 0004), baik berstatus MONOCHROME1 maupun MONOCHROME2."),
        ("3.", "Inversi polaritas citra (inversion check): Mengonversi seluruh citra berstatus MONOCHROME1 menjadi MONOCHROME2 secara konsisten agar polaritas tampilan seragam sebelum proses pemodelan."),
        ("4.", "Pemotongan persentil (percentile clipping): Memangkas nilai intensitas ekstrem pada persentil ke-0,5 dan ke-99,5 untuk mereduksi pengaruh titik pencilan (outliers)."),
        ("5.", "Normalisasi nilai piksel (normalization): Mengonversi rentang intensitas mentah 12/16-bit ke skala ternormalisasi [0, 1] sebagai masukan model DnCNN, serta menyediakan fungsi penskalaan kembali (re-scaling) pada penampil (viewer)."),
        ("6.", "Penyesuaian ukuran (resize/crop): Menyeragamkan ukuran citra dengan mempertahankan aspek rasio piksel asli agar geometri batas organ tetap terjaga."),
        ("7.", "Injeksi derau sintetis (noise injection): Menambahkan simulasi derau campuran Poisson-Gaussian dosis rendah secara terkontrol sebagai data uji eksperimen restorasi.")
    ]

    for num, text in steps:
        p = doc.add_paragraph()
        format_p(p, space_after_pt=1, space_before_pt=0, first_line_indent_pt=14.17)
        add_styled_run(p, num + "\t")
        words = text.split(" ")
        for w_idx, w in enumerate(words):
            sp = " " if w_idx < len(words) - 1 else ""
            clean_w = w.strip("(),:.")
            is_it = clean_w in ["read", "pixel", "array", "photometric", "interpretation", "inversion", "check", "percentile", "clipping", "outliers", "normalization", "re-scaling", "viewer", "resize", "crop", "noise", "injection"]
            add_styled_run(p, w + sp, italic=is_it)

    out_path = r'd:\laragon\www\belajar-laravel\SUBBAB_3_6_PRAPEMROSESAN.docx'
    doc.save(out_path)
    print("Successfully generated ultra-compact Subbab 3.6 docx:", out_path)

if __name__ == '__main__':
    main()
