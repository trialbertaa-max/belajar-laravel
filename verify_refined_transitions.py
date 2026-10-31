import docx

doc = docx.Document('test_transitions_result.docx')

with open('inspect_refined_transitions.txt', 'w', encoding='utf-8') as f:
    def check_context(label, query):
        for i, p in enumerate(doc.paragraphs):
            if query in p.text:
                f.write(f"\n=== Found {label} at p[{i}] ===\n")
                for k in range(max(0, i-2), min(len(doc.paragraphs), i+3)):
                    drawings = doc.paragraphs[k]._p.xpath('.//w:drawing')
                    f.write(f"  p[{k}] (style: {doc.paragraphs[k].style.name}, drawings: {len(drawings)}): {doc.paragraphs[k].text}\n")
                    for r_idx, r in enumerate(doc.paragraphs[k].runs):
                        f.write(f"      run[{r_idx}] i={r.italic} b={r.bold}: {repr(r.text)}\n")

    check_context('Gambar 2.1', 'Gambar 2. 1')
    check_context('Gambar 2.2', 'Gambar 2. 2')
    check_context('Gambar 2.3', 'Gambar 2. 3')
    check_context('Subbab 2.4.4', '2.4.4')

print("Done writing inspect_refined_transitions.txt")
