import docx

doc = docx.Document('d:/laragon/www/belajar-laravel/DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx')

for i, p in enumerate(doc.paragraphs):
    if 'Gambar 2. 2' in p.text or 'Gambar 2. 3' in p.text:
        print(f"p[{i}] text: {p.text}")
        p_prev = doc.paragraphs[i-1]
        print(f"  p[{i-1}] has drawing: {'w:drawing' in p_prev._p.xml}")
