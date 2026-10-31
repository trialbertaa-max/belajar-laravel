import docx

doc = docx.Document('DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx')
print('Total paragraphs in document:', len(doc.paragraphs))

with open('verified_main_transitions.txt', 'w', encoding='utf-8') as f:
    for i, p in enumerate(doc.paragraphs):
        if 'Gambar 2. 1' in p.text:
            f.write(f"\n--- Gambar 2.1 ---\n")
            f.write(f"p[{i-2}] Text: {doc.paragraphs[i-2].text}\n")
            f.write(f"p[{i-1}] Drawings count: {len(doc.paragraphs[i-1]._p.xpath('.//w:drawing'))}\n")
            f.write(f"p[{i}] Caption: {p.text}\n")
        elif 'Gambar 2. 2' in p.text:
            f.write(f"\n--- Gambar 2.2 ---\n")
            f.write(f"p[{i-2}] Text: {doc.paragraphs[i-2].text[-180:]}\n")
            f.write(f"p[{i-1}] Drawings count: {len(doc.paragraphs[i-1]._p.xpath('.//w:drawing'))}\n")
            f.write(f"p[{i}] Caption: {p.text}\n")
        elif 'Gambar 2. 3' in p.text:
            f.write(f"\n--- Gambar 2.3 ---\n")
            f.write(f"p[{i-2}] Text: {doc.paragraphs[i-2].text[-180:]}\n")
            f.write(f"p[{i-1}] Drawings count: {len(doc.paragraphs[i-1]._p.xpath('.//w:drawing'))}\n")
            f.write(f"p[{i}] Caption: {p.text}\n")
        elif '2.4.4' in p.text and p.style.name.startswith('Heading'):
            f.write(f"\n--- Sebelum Heading 2.4.4 ---\n")
            f.write(f"p[{i-1}] Text: {doc.paragraphs[i-1].text[-180:]}\n")
            f.write(f"p[{i}] Heading: {p.text}\n")

print("Done writing verified_main_transitions.txt")
