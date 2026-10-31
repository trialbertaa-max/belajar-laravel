import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, size_pt=12):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt)
    return r

def main():
    doc = docx.Document()

    # Standard Skripsi Margins (Left 4cm, Top/Right/Bottom 3cm)
    for section in doc.sections:
        section.top_margin = Inches(1.18)    # 3 cm
        section.bottom_margin = Inches(1.18) # 3 cm
        section.left_margin = Inches(1.57)   # 4 cm
        section.right_margin = Inches(1.18)  # 3 cm

    # --- JUDUL SUBBAB 3.2 ---
    h2 = doc.add_paragraph()
    format_p(h2, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h2, "3.2\tTempat dan Waktu Penelitian", bold=True, size_pt=12)

    # Paragraf 1: Tempat Penelitian
    p1 = doc.add_paragraph()
    format_p(p1)
    add_styled_run(p1, "Kegiatan penelitian komputasi, pengembangan model restorasi citra, dan perancangan purwarupa stasiun kerja medis dilaksanakan di Laboratorium Komputasi Departemen Fisika, Fakultas Sains dan Teknologi, Universitas Airlangga serta PT Cahaya Medika Healthcare (CMH). Sementara itu, proses kurasi dataset klinis radiografi toraks PA terdeidentifikasi, pengujian transmisi jaringan modalitas rontgen digital, dan validasi klinis acuan baku oleh Dokter Spesialis Radiologi dilaksanakan di Instalasi Radiologi Cahaya Diagnostic Centre (CDC) Surabaya.")

    # Paragraf 2: Waktu Penelitian
    p2 = doc.add_paragraph()
    format_p(p2)
    add_styled_run(p2, "Penelitian ini direncanakan berlangsung selama 6 bulan, terhitung mulai bulan September 2026 hingga Februari 2027. Rincian pembagian matriks jadwal kegiatan penelitian dijabarkan secara terstruktur pada Subbab 3.15.")

    out_path = r'd:\laragon\www\belajar-laravel\SUBBAB_3_2_TEMPAT_DAN_WAKTU.docx'
    doc.save(out_path)
    print("Successfully generated:", out_path)

if __name__ == '__main__':
    main()

