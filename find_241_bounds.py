import docx

doc = docx.Document('d:/laragon/www/belajar-laravel/DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx')

start_idx = None
end_idx = None

for i, p in enumerate(doc.paragraphs):
    if '2.4.1' in p.text:
        start_idx = i
    elif '2.4.2' in p.text:
        end_idx = i
        break

print(f"2.4.1 starts at p[{start_idx}], 2.4.2 starts at p[{end_idx}]")
for i in range(start_idx, end_idx):
    p = doc.paragraphs[i]
    has_math = len(p._p.xpath('.//*[local-name()="oMath"]')) > 0
    print(f"P{i:03d} [math={has_math}]: {p.text[:60]}")
