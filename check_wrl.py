import docx

p = r'C:\Users\Guest User\Downloads\skripsehh\~WRL0003.tmp'
try:
    doc = docx.Document(p)
    print("Opened ~WRL0003.tmp as docx! Total paragraphs:", len(doc.paragraphs))
    for i in range(15):
        print(f"p[{i}]: {doc.paragraphs[i].text[:80]}")
    for i, para in enumerate(doc.paragraphs):
        if '2.4' in para.text:
            print(f"p[{i}]: {para.text[:80]}")
except Exception as e:
    print("Error:", e)
