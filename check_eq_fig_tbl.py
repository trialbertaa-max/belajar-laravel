import docx
import re

doc = docx.Document('DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU.docx')

print("=== CHECKING EQUATIONS ===")
for i, p in enumerate(doc.paragraphs):
    if any(k in p.text.lower() for k in ['persamaan', 'rumus', '(2.', '(3.']):
        print(f"P[{i}]: {p.text.strip()[:100]}")

print("\n=== CHECKING FIGURES ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith('Gambar'):
        print(f"P[{i}] ({p.style.name}): {repr(p.text.strip())}")

print("\n=== CHECKING TABLES ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith('Tabel'):
        print(f"P[{i}] ({p.style.name}): {repr(p.text.strip())}")
