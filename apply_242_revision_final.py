import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.text.paragraph import Paragraph
import lxml.etree as ET
import shutil
import copy
import os

# Load XSLT
xslt_path = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
xslt = ET.parse(xslt_path)
transform = ET.XSLT(xslt)

def mml_to_omml_element(mml_str):
    mml_tree = ET.fromstring(mml_str)
    omml_tree = transform(mml_tree)
    root = omml_tree.getroot()
    if root.tag.endswith('oMath'):
        para = ET.Element('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara')
        para_pr = ET.SubElement(para, '{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathParaPr')
        jc = ET.SubElement(para_pr, '{http://schemas.openxmlformats.org/officeDocument/2006/math}jc')
        jc.set('{http://schemas.openxmlformats.org/officeDocument/2006/math}val', 'center')
        para.append(root)
        root = para
    xml_bytes = ET.tostring(root, encoding='utf-8')
    return parse_xml(xml_bytes)

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False, subscript=False, superscript=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic
    if subscript:
        r.font.subscript = True
    if superscript:
        r.font.superscript = True

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, subscript=False, superscript=False, size_pt=12):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, subscript=subscript, superscript=superscript, size_pt=size_pt)
    return r

def build_section_242(target_before_p, drawing_fig2_2):
    def create_p(style_name='Normal'):
        new_p_elm = OxmlElement('w:p')
        target_before_p._p.addprevious(new_p_elm)
        p = Paragraph(new_p_elm, target_before_p._parent)
        if style_name != 'Normal':
            p.style = style_name
        return p

    def create_eq_p(mml_str):
        p = create_p()
        format_p(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, line_spacing=1.0, first_line_indent_pt=0)
        omml_elm = mml_to_omml_element(mml_str)
        p._p.append(omml_elm)
        return p

    # 1. Heading 3: 2.4.2 Bilateral Filter
    h3 = create_p('Heading 3')
    format_p(h3, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=4, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h3, "2.4.2\t", bold=True, italic=False)
    add_styled_run(h3, "Bilateral Filter", bold=True, italic=True)

    # 2. Paragraf 1: Definisi & Prinsip Dasar
    p1 = create_p()
    format_p(p1)
    add_styled_run(p1, "Bilateral filter", italic=True)
    add_styled_run(p1, " merupakan metode penapisan ranah spasial non-linear (")
    add_styled_run(p1, "non-linear spatial domain filter", italic=True)
    add_styled_run(p1, ") yang dirancang khusus untuk mereduksi derau sembari mempertahankan ketajaman garis batas tepi objek (")
    add_styled_run(p1, "edge-preserving smoothing", italic=True)
    add_styled_run(p1, ") (Tomasi & Manduchi, 1998). Berbeda dengan ")
    add_styled_run(p1, "Gaussian filter", italic=True)
    add_styled_run(p1, " yang hanya memperhitungkan jarak kedekatan geometris antar-piksel, ")
    add_styled_run(p1, "Bilateral filter", italic=True)
    add_styled_run(p1, " menggabungkan dua fungsi pembobot ")
    add_styled_run(p1, "Gaussian", italic=True)
    add_styled_run(p1, " secara simultan: fungsi ranah spasial (")
    add_styled_run(p1, "G", italic=True)
    add_styled_run(p1, "σs", italic=True, subscript=True)
    add_styled_run(p1, ") yang mengukur kedekatan posisi geometris, dan fungsi rentang fotometrik radiometrik (")
    add_styled_run(p1, "G", italic=True)
    add_styled_run(p1, "σr", italic=True, subscript=True)
    add_styled_run(p1, ") yang mengukur kemiripan nilai derajat keabuan (")
    add_styled_run(p1, "intensity similarity", italic=True)
    add_styled_run(p1, "). Nilai intensitas piksel terestorasi ")
    add_styled_run(p1, "I", italic=True)
    add_styled_run(p1, "bilateral", italic=False, subscript=True)
    add_styled_run(p1, "(")
    add_styled_run(p1, "p", italic=True)
    add_styled_run(p1, ") pada koordinat target ")
    add_styled_run(p1, "p", italic=True)
    add_styled_run(p1, " diformulasikan melalui persamaan matematis:")

    # 3. Persamaan 1: Formulasi Utama Bilateral Filter
    mml1 = '''<math xmlns="http://www.w3.org/1998/Math/MathML">
      <msub><mi>I</mi><mtext>bilateral</mtext></msub>
      <mo>(</mo><mi>p</mi><mo>)</mo>
      <mo>=</mo>
      <mfrac>
        <mn>1</mn>
        <msub><mi>W</mi><mi>p</mi></msub>
      </mfrac>
      <munder>
        <mo>∑</mo>
        <mrow><mi>q</mi><mo>∈</mo><mi>Ω</mi></mrow>
      </munder>
      <msub><mi>G</mi><msub><mi>σ</mi><mi>s</mi></msub></msub>
      <mo>(</mo><mo>‖</mo><mi>p</mi><mo>-</mo><mi>q</mi><mo>‖</mo><mo>)</mo>
      <mo>·</mo>
      <msub><mi>G</mi><msub><mi>σ</mi><mi>r</mi></msub></msub>
      <mo>(</mo><mo>|</mo><mi>I</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>-</mo><mi>I</mi><mo>(</mo><mi>q</mi><mo>)</mo><mo>|</mo><mo>)</mo>
      <mo>·</mo>
      <mi>I</mi>
      <mo>(</mo><mi>q</mi><mo>)</mo>
    </math>'''
    create_eq_p(mml1)

    # 4. Persamaan 2: Normalisasi W_p
    mml2 = '''<math xmlns="http://www.w3.org/1998/Math/MathML">
      <msub><mi>W</mi><mi>p</mi></msub>
      <mo>=</mo>
      <munder>
        <mo>∑</mo>
        <mrow><mi>q</mi><mo>∈</mo><mi>Ω</mi></mrow>
      </munder>
      <msub><mi>G</mi><msub><mi>σ</mi><mi>s</mi></msub></msub>
      <mo>(</mo><mo>‖</mo><mi>p</mi><mo>-</mo><mi>q</mi><mo>‖</mo><mo>)</mo>
      <mo>·</mo>
      <msub><mi>G</mi><msub><mi>σ</mi><mi>r</mi></msub></msub>
      <mo>(</mo><mo>|</mo><mi>I</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>-</mo><mi>I</mi><mo>(</mo><mi>q</mi><mo>)</mo><mo>|</mo><mo>)</mo>
    </math>'''
    create_eq_p(mml2)

    # 5. Paragraf 2: Keterangan Variabel & Pengantar Parameter
    p2 = create_p()
    format_p(p2)
    add_styled_run(p2, "Keterangan parameter matematis: ")
    add_styled_run(p2, "Ω", italic=True)
    add_styled_run(p2, " merepresentasikan jendela ketetanggaan spasial lokal (")
    add_styled_run(p2, "spatial neighborhood window", italic=True)
    add_styled_run(p2, ") yang berpusat pada piksel target ")
    add_styled_run(p2, "p", italic=True)
    add_styled_run(p2, ", ||")
    add_styled_run(p2, "p", italic=True)
    add_styled_run(p2, " \u2212 ")
    add_styled_run(p2, "q", italic=True)
    add_styled_run(p2, "|| menyatakan jarak ")
    add_styled_run(p2, "Euclidean", italic=True)
    add_styled_run(p2, " antar-piksel dalam bidang koordinat citra, dan |")
    add_styled_run(p2, "I", italic=True)
    add_styled_run(p2, "(")
    add_styled_run(p2, "p", italic=True)
    add_styled_run(p2, ") \u2212 ")
    add_styled_run(p2, "I", italic=True)
    add_styled_run(p2, "(")
    add_styled_run(p2, "q", italic=True)
    add_styled_run(p2, ")| menyatakan selisih mutlak nilai intensitas fotometrik antara piksel ")
    add_styled_run(p2, "p", italic=True)
    add_styled_run(p2, " dan piksel tetangga ")
    add_styled_run(p2, "q", italic=True)
    add_styled_run(p2, ". Faktor pembagi ")
    add_styled_run(p2, "W", italic=True)
    add_styled_run(p2, "p", italic=True, subscript=True)
    add_styled_run(p2, " bertindak sebagai faktor normalisasi bobot untuk menjamin kelestarian energi rerata intensitas (")
    add_styled_run(p2, "unit gain", italic=True)
    add_styled_run(p2, "). Kinerja penapisan ")
    add_styled_run(p2, "Bilateral filter", italic=True)
    add_styled_run(p2, " dikendalikan secara presisi oleh interaksi dua parameter independen melalui formulasi ")
    add_styled_run(p2, "kernel", italic=True)
    add_styled_run(p2, " penyusunnya:")

    # 6. Poin a: Parameter Spasial sigma_s
    pa = create_p()
    format_p(pa)
    add_styled_run(pa, "a.\t", bold=True, italic=False)
    add_styled_run(pa, "Parameter Spasial (", bold=True, italic=False)
    add_styled_run(pa, "σ", bold=True, italic=True)
    add_styled_run(pa, "s", bold=True, italic=True, subscript=True)
    add_styled_run(pa, "): ", bold=True, italic=False)
    add_styled_run(pa, "Fungsi pembobot ranah spasial ")
    add_styled_run(pa, "G", italic=True)
    add_styled_run(pa, "σs", italic=True, subscript=True)
    add_styled_run(pa, " mengukur kedekatan geometris antar-piksel melalui formulasi matematis:")

    # Persamaan 3: G_sigma_s
    mml3 = '''<math xmlns="http://www.w3.org/1998/Math/MathML">
      <msub><mi>G</mi><msub><mi>σ</mi><mi>s</mi></msub></msub>
      <mo>(</mo><mo>‖</mo><mi>p</mi><mo>-</mo><mi>q</mi><mo>‖</mo><mo>)</mo>
      <mo>=</mo>
      <mi>exp</mi>
      <mfenced>
        <mrow>
          <mo>-</mo>
          <mfrac>
            <msup><mrow><mo>‖</mo><mi>p</mi><mo>-</mo><mi>q</mi><mo>‖</mo></mrow><mn>2</mn></msup>
            <mrow><mn>2</mn><msubsup><mi>σ</mi><mi>s</mi><mn>2</mn></msubsup></mrow>
          </mfrac>
        </mrow>
      </mfenced>
    </math>'''
    create_eq_p(mml3)

    pa_desc = create_p()
    format_p(pa_desc)
    add_styled_run(pa_desc, "Parameter ")
    add_styled_run(pa_desc, "σ", italic=True)
    add_styled_run(pa_desc, "s", italic=True, subscript=True)
    add_styled_run(pa_desc, " mengendalikan skala kedekatan spasial (")
    add_styled_run(pa_desc, "spatial proximity scale", italic=True)
    add_styled_run(pa_desc, ") yang menentukan radius ketetanggaan piksel di dalam jendela ")
    add_styled_run(pa_desc, "Ω", italic=True)
    add_styled_run(pa_desc, " yang dilibatkan dalam perhitungan rata-rata terbobot. Pemilihan nilai ")
    add_styled_run(pa_desc, "σ", italic=True)
    add_styled_run(pa_desc, "s", italic=True, subscript=True)
    add_styled_run(pa_desc, " yang moderat diperlukan pada radiografi toraks untuk menangkap korelasi spasial fluktuasi derau bintik kuantum (")
    add_styled_run(pa_desc, "quantum mottle", italic=True)
    add_styled_run(pa_desc, ") tanpa memperluas jendela komputasi secara berlebihan.")

    # 7. Poin b: Parameter Rentang Intensitas sigma_r
    pb = create_p()
    format_p(pb)
    add_styled_run(pb, "b.\t", bold=True, italic=False)
    add_styled_run(pb, "Parameter Rentang Intensitas (", bold=True, italic=False)
    add_styled_run(pb, "σ", bold=True, italic=True)
    add_styled_run(pb, "r", bold=True, italic=True, subscript=True)
    add_styled_run(pb, "): ", bold=True, italic=False)
    add_styled_run(pb, "Fungsi pembobot rentang radiometrik ")
    add_styled_run(pb, "G", italic=True)
    add_styled_run(pb, "σr", italic=True, subscript=True)
    add_styled_run(pb, " mengukur kemiripan nilai derajat keabuan fotometrik melalui formulasi matematis:")

    # Persamaan 4: G_sigma_r
    mml4 = '''<math xmlns="http://www.w3.org/1998/Math/MathML">
      <msub><mi>G</mi><msub><mi>σ</mi><mi>r</mi></msub></msub>
      <mo>(</mo><mo>|</mo><mi>I</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>-</mo><mi>I</mi><mo>(</mo><mi>q</mi><mo>)</mo><mo>|</mo><mo>)</mo>
      <mo>=</mo>
      <mi>exp</mi>
      <mfenced>
        <mrow>
          <mo>-</mo>
          <mfrac>
            <msup><mrow><mo>|</mo><mi>I</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>-</mo><mi>I</mi><mo>(</mo><mi>q</mi><mo>)</mo><mo>|</mo></mrow><mn>2</mn></msup>
            <mrow><mn>2</mn><msubsup><mi>σ</mi><mi>r</mi><mn>2</mn></msubsup></mrow>
          </mfrac>
        </mrow>
      </mfenced>
    </math>'''
    create_eq_p(mml4)

    pb_desc = create_p()
    format_p(pb_desc)
    add_styled_run(pb_desc, "Parameter ")
    add_styled_run(pb_desc, "σ", italic=True)
    add_styled_run(pb_desc, "r", italic=True, subscript=True)
    add_styled_run(pb_desc, " bertindak sebagai ambang batas fotometrik (")
    add_styled_run(pb_desc, "photometric tonal threshold", italic=True)
    add_styled_run(pb_desc, ") yang mengendalikan sensitivitas penapis terhadap perbedaan kontras radiometrik antar-piksel. Parameter ini menentukan batasan seberapa besar selisih intensitas yang masih dianggap sebagai fluktuasi derau acak atau telah tergolong sebagai batas tepi anatomis (")
    add_styled_run(pb_desc, "anatomical boundary", italic=True)
    add_styled_run(pb_desc, ").")

    # 8. Paragraf 3: Dinamika Adaptif pada Jaringan Toraks & Pelestarian Batas Jantung
    p3 = create_p()
    format_p(p3)
    add_styled_run(p3, "Dinamika adaptif ")
    add_styled_run(p3, "Bilateral filter", italic=True)
    add_styled_run(p3, " pada citra radiografi toraks proyeksi ")
    add_styled_run(p3, "Posteroanterior", italic=True)
    add_styled_run(p3, " (PA) bertumpu pada interaksi kedua fungsi pembobot tersebut terhadap variasi koefisien atenuasi radiologis. Pada area lapangan paru (")
    add_styled_run(p3, "lung field", italic=True)
    add_styled_run(p3, ") yang relatif homogen, selisih intensitas antar-piksel tetangga bernilai sangat rendah (|")
    add_styled_run(p3, "I", italic=True)
    add_styled_run(p3, "(")
    add_styled_run(p3, "p", italic=True)
    add_styled_run(p3, ") \u2212 ")
    add_styled_run(p3, "I", italic=True)
    add_styled_run(p3, "(")
    add_styled_run(p3, "q", italic=True)
    add_styled_run(p3, ")| \u2248 0), sehingga bobot rentang bernilai mendekati satu (")
    add_styled_run(p3, "G", italic=True)
    add_styled_run(p3, "σr", italic=True, subscript=True)
    add_styled_run(p3, " \u2248 1). Dalam kondisi tersebut, ")
    add_styled_run(p3, "Bilateral filter", italic=True)
    add_styled_run(p3, " beroperasi menyerupai ")
    add_styled_run(p3, "Gaussian filter", italic=True)
    add_styled_run(p3, " standar yang meratakan derau kuantum secara maksimal melalui perataan lokal (")
    add_styled_run(p3, "local averaging", italic=True)
    add_styled_run(p3, "). Sebaliknya, ketika jendela penapisan melintasi batas transisi anatomis kardioperikardial antara bayangan siluet kardiak yang radioopak (densitas tinggi) dan parenkim paru yang radiolusen (densitas rendah), selisih intensitas melonjak tajam melampaui nilai ambang (|")
    add_styled_run(p3, "I", italic=True)
    add_styled_run(p3, "(")
    add_styled_run(p3, "p", italic=True)
    add_styled_run(p3, ") \u2212 ")
    add_styled_run(p3, "I", italic=True)
    add_styled_run(p3, "(")
    add_styled_run(p3, "q", italic=True)
    add_styled_run(p3, ")| >> ")
    add_styled_run(p3, "σ", italic=True)
    add_styled_run(p3, "r", italic=True, subscript=True)
    add_styled_run(p3, "). Lonjakan selisih ini menyebabkan bobot rentang seketika anjlok mendekati nol (")
    add_styled_run(p3, "G", italic=True)
    add_styled_run(p3, "σr", italic=True, subscript=True)
    add_styled_run(p3, " \u2192 0), sehingga kontribusi intensitas piksel parenkim paru diabaikan dari perhitungan siluet kardiak. Mekanisme pelestarian tepi (")
    add_styled_run(p3, "edge preservation", italic=True)
    add_styled_run(p3, ") ini menjaga kecuraman gradien garis batas lateral jantung (")
    add_styled_run(p3, "maximum right diameter", italic=True)
    add_styled_run(p3, " dan ")
    add_styled_run(p3, "maximum left diameter", italic=True)
    add_styled_run(p3, ") serta ketajaman sudut sinus kostofrenikus (")
    add_styled_run(p3, "costophrenic angle", italic=True)
    add_styled_run(p3, ") tetap terlindungi dari efek pengaburan tepi (")
    add_styled_run(p3, "edge blurring", italic=True)
    add_styled_run(p3, ").")

    # 9. Paragraf 4: Analisis Teoretis Trade-off, Risiko Over-smoothing & Staircasing, serta Grid Search
    p4 = create_p()
    format_p(p4)
    add_styled_run(p4, "Meskipun memiliki keunggulan dalam melestarikan batas tepi objek, efektivitas restorasi ")
    add_styled_run(p4, "Bilateral filter", italic=True)
    add_styled_run(p4, " sangat bergantung pada ketepatan pemilihan pasangan parameter (")
    add_styled_run(p4, "σ", italic=True)
    add_styled_run(p4, "s", italic=True, subscript=True)
    add_styled_run(p4, ", ")
    add_styled_run(p4, "σ", italic=True)
    add_styled_run(p4, "r", italic=True, subscript=True)
    add_styled_run(p4, "). Apabila parameter rentang ")
    add_styled_run(p4, "σ", italic=True)
    add_styled_run(p4, "r", italic=True, subscript=True)
    add_styled_run(p4, " ditetapkan terlalu besar (")
    add_styled_run(p4, "σ", italic=True)
    add_styled_run(p4, "r", italic=True, subscript=True)
    add_styled_run(p4, " \u2192 \u221e), gerbang ambang batas fotometrik kehilangan kemampuan selektivitasnya (")
    add_styled_run(p4, "G", italic=True)
    add_styled_run(p4, "σr", italic=True, subscript=True)
    add_styled_run(p4, " \u2192 1 untuk seluruh piksel), sehingga ")
    add_styled_run(p4, "Bilateral filter", italic=True)
    add_styled_run(p4, " berdegenerasi menjadi penapis ")
    add_styled_run(p4, "isotropic Gaussian filter", italic=True)
    add_styled_run(p4, ". Kondisi ini memicu fenomena pengaburan berlebih (")
    add_styled_run(p4, "over-smoothing", italic=True)
    add_styled_run(p4, ") yang mengikis ketajaman batas anatomis lateral jantung dan memperbesar galat klinis pengukuran ")
    add_styled_run(p4, "Cardiothoracic Ratio", italic=True)
    add_styled_run(p4, " (CTR). Sebaliknya, apabila ")
    add_styled_run(p4, "σ", italic=True)
    add_styled_run(p4, "r", italic=True, subscript=True)
    add_styled_run(p4, " diatur terlalu kecil (")
    add_styled_run(p4, "σ", italic=True)
    add_styled_run(p4, "r", italic=True, subscript=True)
    add_styled_run(p4, " \u2192 0), penapis menjadi terlampau restriktif sehingga penekanan derau tidak optimal dan berisiko memunculkan artefak efek tangga (")
    add_styled_run(p4, "staircasing artifact", italic=True)
    add_styled_run(p4, " atau pembentukan pulau intensitas palsu). Selain itu, nilai ")
    add_styled_run(p4, "σ", italic=True)
    add_styled_run(p4, "s", italic=True, subscript=True)
    add_styled_run(p4, " yang berlebihan akan memperluas jendela komputasi melampaui struktur toraks lokal dan meningkatkan beban komputasi secara kuadratik. Oleh karena itu, penentuan kombinasi parameter (")
    add_styled_run(p4, "σ", italic=True)
    add_styled_run(p4, "s", italic=True, subscript=True)
    add_styled_run(p4, ", ")
    add_styled_run(p4, "σ", italic=True)
    add_styled_run(p4, "r", italic=True, subscript=True)
    add_styled_run(p4, ") secara sistematik melalui penelusuran kisi (")
    add_styled_run(p4, "grid search", italic=True)
    add_styled_run(p4, ") pada set validasi menjadi tahapan optimasi yang krusial sebelum citra hasil restorasi dievaluasi menggunakan metrik berbasis tugas klinis (")
    add_styled_run(p4, "task-based evaluation", italic=True)
    add_styled_run(p4, ").")

    # 10. Paragraf 5: Pengantar Gambar 2.2
    p5 = create_p()
    format_p(p5)
    add_styled_run(p5, "Prinsip kerja penapisan ")
    add_styled_run(p5, "bilateral", italic=True)
    add_styled_run(p5, " yang memadukan bobot kedekatan geometris ranah spasial dan bobot kemiripan radiometrik dalam mempertahankan ketajaman garis batas tepi kardiak diilustrasikan secara komprehensif pada Gambar 2.2.")

    # 11. Gambar 2.2 Paragraph
    p_fig = create_p()
    format_p(p_fig, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, line_spacing=1.0, first_line_indent_pt=0)
    drawing_copy = copy.deepcopy(drawing_fig2_2)
    run_fig = p_fig.add_run()
    run_fig._r.append(drawing_copy)

    # 12. Caption Gambar 2.2 Paragraph
    p_cap = create_p('Caption')
    format_p(p_cap, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(p_cap, "Gambar 2. 2 Prinsip Kerja ", bold=False, italic=False)
    add_styled_run(p_cap, "Bilateral Filtering", bold=False, italic=True)
    add_styled_run(p_cap, " dalam Mereduksi Derau sembari Mempertahankan Batas Tepi Organ (Tomasi & Manduchi, 1998)", bold=False, italic=False)

def main():
    doc_path = r'd:\laragon\www\belajar-laravel\DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx'
    backup_path = r'd:\laragon\www\belajar-laravel\DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP_PRE_242_REVISI.docx'

    # 1. Create safety backup
    shutil.copy2(doc_path, backup_path)
    print(f"Safety backup saved to {backup_path}")

    # 2. Open document
    doc = docx.Document(doc_path)
    start_idx = None
    end_idx = None
    drawing_fig2_2 = None

    for i, p in enumerate(doc.paragraphs):
        if '2.4.2' in p.text:
            start_idx = i
        elif 'Gambar 2. 2' in p.text:
            prev_p = doc.paragraphs[i-1]
            drawings = prev_p._p.xpath('.//w:drawing')
            if drawings:
                drawing_fig2_2 = copy.deepcopy(drawings[0])
        elif '2.4.3' in p.text:
            end_idx = i
            break

    print(f"Original 2.4.2 range: p[{start_idx}] to p[{end_idx}]")
    print(f"Extracted drawing_fig2_2: {drawing_fig2_2 is not None}")

    target_before = doc.paragraphs[end_idx]

    # 3. Insert new 2.4.2 paragraphs
    build_section_242(target_before, drawing_fig2_2)

    # 4. Remove old 2.4.2 paragraphs
    old_p_elms = [doc.paragraphs[k]._p for k in range(start_idx, end_idx)]
    for elm in old_p_elms:
        elm.getparent().remove(elm)

    # 5. Save back to primary document
    doc.save(doc_path)
    print(f"Successfully updated Subbab 2.4.2 in {doc_path}")

if __name__ == '__main__':
    main()
