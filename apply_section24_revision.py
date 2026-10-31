import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.text.paragraph import Paragraph
import lxml.etree as ET
import shutil
import copy
import os

# Load XSLT for MathML to OMML
xslt_path = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
xslt = ET.parse(xslt_path)
transform = ET.XSLT(xslt)

def mml_to_omml_element(mml_str):
    """Converts MathML string to an OpenXML m:oMathPara element."""
    mml_tree = ET.fromstring(mml_str)
    omml_tree = transform(mml_tree)
    root = omml_tree.getroot()
    if root.tag.endswith('oMath'):
        para = ET.Element('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara')
        para_pr = ET.SubElement(para, '{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathParaPr')
        jc = ET.SubElement(para_pr, '{http://schemas.openxmlformats.org/officeDocument/2006/math}jc')
        jc.set('{http://schemas.openxmlformats.org/officeDocument/2006/math}val', 'center')
        para.append(root)
        root = para
    xml_bytes = ET.tostring(root, encoding='utf-8')
    return parse_xml(xml_bytes)

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, size_pt=12):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt)
    return r

def run():
    doc_path = r'd:\laragon\www\belajar-laravel\DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx'
    backup_path = r'd:\laragon\www\belajar-laravel\DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP_PRE_24.docx'
    
    # Restore from pristine backup if needed
    if os.path.exists(backup_path):
        shutil.copy2(backup_path, doc_path)
    else:
        shutil.copy2(doc_path, backup_path)
    print(f"Working on {doc_path}")

    doc = docx.Document(doc_path)
    print(f"Total initial paragraphs: {len(doc.paragraphs)}")

    start_idx = None
    end_idx = None
    drawing_fig2_2 = None
    drawing_fig2_3 = None

    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if 'Metode Restorasi Citra Medis' in t:
            start_idx = i
        elif 'Gambar 2. 2' in t:
            prev_p = doc.paragraphs[i-1]
            drawings = prev_p._p.xpath('.//w:drawing')
            if drawings:
                drawing_fig2_2 = copy.deepcopy(drawings[0])
        elif 'Gambar 2. 3' in t:
            prev_p = doc.paragraphs[i-1]
            drawings = prev_p._p.xpath('.//w:drawing')
            if drawings:
                drawing_fig2_3 = copy.deepcopy(drawings[0])
        elif start_idx is not None and 'Metrik Evaluasi Kualitas Citra' in t:
            end_idx = i
            break

    print(f"Section 2.4 range: p[{start_idx}] to p[{end_idx}]")
    print(f"Preserved Gambar 2.2 drawing: {drawing_fig2_2 is not None}")
    print(f"Preserved Gambar 2.3 drawing: {drawing_fig2_3 is not None}")

    target_p = doc.paragraphs[end_idx]

    def create_p():
        new_p_elm = OxmlElement('w:p')
        target_p._p.addprevious(new_p_elm)
        return Paragraph(new_p_elm, target_p._parent)

    def create_eq_p(mml_str):
        p = create_p()
        format_p(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, line_spacing=1.0, first_line_indent_pt=0)
        omml_elm = mml_to_omml_element(mml_str)
        p._p.append(omml_elm)
        return p

    # =========================================================================
    # HEADING 2: 2.4 Restorasi Citra Digital Radiografi
    # =========================================================================
    h2 = create_p()
    format_p(h2, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h2, "2.4\tRestorasi Citra Digital Radiografi", bold=True, size_pt=12)

    # Paragraf 1
    p = create_p()
    format_p(p)
    add_styled_run(p, "Restorasi citra medis (")
    add_styled_run(p, "medical image restoration", italic=True)
    add_styled_run(p, ") merupakan proses rekonstruksi untuk mengestimasi citra anatomis asli yang bersih dari derau (")
    add_styled_run(p, "clean signal", italic=True)
    add_styled_run(p, ") berdasarkan pemodelan fisis kebalikan (")
    add_styled_run(p, "inverse problem", italic=True)
    add_styled_run(p, ") dari proses pembentukan degradasi (Gonzalez & Woods, 2018). Berbeda dengan peningkatan kualitas citra (")
    add_styled_run(p, "image enhancement", italic=True)
    add_styled_run(p, ") yang bersifat heuristik dan subjektif untuk kenyamanan visual pengamat (seperti penyesuaian kontras lokal), restorasi citra berlandaskan pada pemulihan kuantitatif nilai intensitas fotometrik agar mendekati kondisi transmisi radiasi tanpa derau.")

    # Paragraf 2
    p = create_p()
    format_p(p)
    add_styled_run(p, "Dalam konteks radiografi toraks proyeksi Posteroanterior (PA), keberhasilan metode restorasi tidak hanya diukur berdasarkan peningkatan rasio sinyal terhadap derau secara umum, melainkan pada kemampuan mempertahankan ketajaman garis batas tepi organ (")
    add_styled_run(p, "edge-preserving capability", italic=True)
    add_styled_run(p, "). Garis batas lateral kardioperikardial kanan-kiri dan lengkung internal dinding kosta lateral merupakan batas anatomis kritis yang menentukan keakuratan penentuan diameter jantung (")
    add_styled_run(p, "DH", italic=True)
    add_styled_run(p, ") dan diameter toraks (")
    add_styled_run(p, "DT", italic=True)
    add_styled_run(p, ") pada pengukuran ")
    add_styled_run(p, "Cardiothoracic Ratio", italic=True)
    add_styled_run(p, " (CTR). Berdasarkan mekanisme matematis dan arsitekturnya, metode restorasi dalam penelitian ini diklasifikasikan ke dalam tiga pendekatan: filter linear konvensional (Gaussian Filter), filter non-linear berbasis penapis tepi lokal (Bilateral Filter), dan jaringan saraf tiruan konvolusional dalam berbasis pembelajaran residual (DnCNN).")

    # =========================================================================
    # HEADING 3: 2.4.1 Gaussian Filter
    # =========================================================================
    h3_1 = create_p()
    format_p(h3_1, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=4, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h3_1, "2.4.1\tGaussian Filter", bold=True, size_pt=12)

    # Paragraf 1 Gaussian
    p = create_p()
    format_p(p)
    add_styled_run(p, "Gaussian Filter merupakan filter linear spasial domain yang bekerja melalui operasi konvolusi diskret dua dimensi menggunakan fungsi pembobotan terdistribusi normal (Gaussian) (Gonzalez & Woods, 2018). Fungsi pembobot kernel Gaussian kontinu dua dimensi didefinisikan melalui persamaan matematis:")

    # Eq 1: Gaussian kernel
    create_eq_p('<math xmlns="http://www.w3.org/1998/Math/MathML"><msub><mi>G</mi><mi>σ</mi></msub><mo>(</mo><mi>x</mi><mo>,</mo><mi>y</mi><mo>)</mo><mo>=</mo><mfrac><mn>1</mn><mrow><mn>2</mn><mi>π</mi><msup><mi>σ</mi><mn>2</mn></msup></mrow></mfrac><mi>exp</mi><mfenced><mrow><mo>-</mo><mfrac><mrow><msup><mi>x</mi><mn>2</mn></msup><mo>+</mo><msup><mi>y</mi><mn>2</mn></msup></mrow><mrow><mn>2</mn><msup><mi>σ</mi><mn>2</mn></msup></mrow></mfrac></mrow></mfenced></math>')

    # Paragraf 2 Gaussian
    p = create_p()
    format_p(p)
    add_styled_run(p, "Operasi penyaringan pada piksel target berkoordinat ")
    add_styled_run(p, "p = (x, y)", italic=True)
    add_styled_run(p, " dihitung dengan menjumlahkan perkalian intensitas piksel tetangga ")
    add_styled_run(p, "q \u2208 \u03a9", italic=True)
    add_styled_run(p, " terhadap bobot jarak spasialnya:")

    # Eq 2: Gaussian convolution
    create_eq_p('<math xmlns="http://www.w3.org/1998/Math/MathML"><msub><mi>I</mi><mtext>Gaussian</mtext></msub><mo>(</mo><mi>p</mi><mo>)</mo><mo>=</mo><munder><mo>∑</mo><mrow><mi>q</mi><mo>∈</mo><mi>Ω</mi></mrow></munder><msub><mi>G</mi><mi>σ</mi></msub><mo>(</mo><mo>‖</mo><mi>p</mi><mo>-</mo><mi>q</mi><mo>‖</mo><mo>)</mo><mo>·</mo><mi>I</mi><mo>(</mo><mi>q</mi><mo>)</mo></math>')

    # Paragraf 3 Gaussian
    p = create_p()
    format_p(p)
    add_styled_run(p, "di mana ")
    add_styled_run(p, "\u03a9", italic=True)
    add_styled_run(p, " merepresentasikan jendela lingkungan spasial lokal di sekitar piksel ")
    add_styled_run(p, "p", italic=True)
    add_styled_run(p, ", ")
    add_styled_run(p, "||p \u2212 q||", italic=True)
    add_styled_run(p, " adalah jarak Euclidean antar-piksel, dan parameter ")
    add_styled_run(p, "\u03c3", italic=True)
    add_styled_run(p, " merupakan simpangan baku spasial yang mengatur lebar sebaran fungsi pembobot. Secara fisis, Gaussian filter bekerja sangat efektif meredam fluktuasi derau acak frekuensi tinggi pada area homogen melalui perataan rata-rata lokal (")
    add_styled_run(p, "local averaging", italic=True)
    add_styled_run(p, "). Namun, kelemahan mendasar Gaussian filter terletak pada sifatnya yang ")
    add_styled_run(p, "isotropic", italic=True)
    add_styled_run(p, " dan tidak memperhitungkan informasi variasi nilai fotometrik antar-piksel.")

    # Paragraf 4 Gaussian
    p = create_p()
    format_p(p)
    add_styled_run(p, "Ketika kernel konvolusi Gaussian melintasi perbatasan antara bayangan jantung yang radioopak dan parenkim paru yang radiolusen, nilai intensitas dari kedua kompartemen jaringan tersebut akan diratakan secara merata. Hal ini memicu efek pengaburan tepi (")
    add_styled_run(p, "edge blurring", italic=True)
    add_styled_run(p, ") yang parah, melandaikan gradien transisi tepi kardioperikardial, serta mengikis ketajaman garis batas sinus kostofrenikus dan dinding kosta. Akibatnya, ketidakpastian spasial dalam penentuan titik terluar batas kardiak meningkat secara signifikan. Dalam penelitian ini, Gaussian filter diikutsertakan sebagai metode ")
    add_styled_run(p, "baseline", italic=True)
    add_styled_run(p, " pembanding klasik untuk membuktikan secara empiris bahwa pereduksian derau tanpa kemampuan pelestarian tepi justru memperburuk galat pengukuran CTR.")

    # =========================================================================
    # HEADING 3: 2.4.2 Bilateral Filter
    # =========================================================================
    h3_2 = create_p()
    format_p(h3_2, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=4, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h3_2, "2.4.2\tBilateral Filter", bold=True, size_pt=12)

    # Paragraf 1 Bilateral
    p = create_p()
    format_p(p)
    add_styled_run(p, "Bilateral Filter merupakan filter spasial non-linear yang dirancang khusus untuk menghaluskan derau sembari mempertahankan ketajaman garis batas tepi objek (")
    add_styled_run(p, "edge-preserving smoothing", italic=True)
    add_styled_run(p, ") (Tomasi & Manduchi, 1998). Berbeda dengan filter konvensional yang hanya memperhitungkan kedekatan jarak geometris, Bilateral Filter menggabungkan dua fungsi pembobot Gaussian secara simultan: fungsi domain spasial (")
    add_styled_run(p, "G\u03c3s", italic=True)
    add_styled_run(p, ") yang mengukur kedekatan posisi geometris, dan fungsi rentang radiometrik (")
    add_styled_run(p, "G\u03c3r", italic=True)
    add_styled_run(p, ") yang mengukur kemiripan nilai intensitas piksel. Intensitas piksel terestorasi pada titik ")
    add_styled_run(p, "p", italic=True)
    add_styled_run(p, " diformulasikan melalui persamaan matematis:")

    # Eq 3: Bilateral Ifiltered
    create_eq_p('<math xmlns="http://www.w3.org/1998/Math/MathML"><msub><mi>I</mi><mtext>filtered</mtext></msub><mo>(</mo><mi>p</mi><mo>)</mo><mo>=</mo><mfrac><mn>1</mn><msub><mi>W</mi><mi>p</mi></msub></mfrac><munder><mo>∑</mo><mrow><mi>q</mi><mo>∈</mo><mi>Ω</mi></mrow></munder><msub><mi>G</mi><msub><mi>σ</mi><mi>s</mi></msub></msub><mo>(</mo><mo>‖</mo><mi>p</mi><mo>-</mo><mi>q</mi><mo>‖</mo><mo>)</mo><mo>·</mo><msub><mi>G</mi><msub><mi>σ</mi><mi>r</mi></msub></msub><mo>(</mo><mo>|</mo><mi>I</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>-</mo><mi>I</mi><mo>(</mo><mi>q</mi><mo>)</mo><mo>|</mo><mo>)</mo><mo>·</mo><mi>I</mi><mo>(</mo><mi>q</mi><mo>)</mo></math>')

    # Eq 4: Bilateral Wp
    create_eq_p('<math xmlns="http://www.w3.org/1998/Math/MathML"><msub><mi>W</mi><mi>p</mi></msub><mo>=</mo><munder><mo>∑</mo><mrow><mi>q</mi><mo>∈</mo><mi>Ω</mi></mrow></munder><msub><mi>G</mi><msub><mi>σ</mi><mi>s</mi></msub></msub><mo>(</mo><mo>‖</mo><mi>p</mi><mo>-</mo><mi>q</mi><mo>‖</mo><mo>)</mo><mo>·</mo><msub><mi>G</mi><msub><mi>σ</mi><mi>r</mi></msub></msub><mo>(</mo><mo>|</mo><mi>I</mi><mo>(</mo><mi>p</mi><mo>)</mo><mo>-</mo><mi>I</mi><mo>(</mo><mi>q</mi><mo>)</mo><mo>|</mo><mo>)</mo></math>')

    # Paragraf 2 Bilateral (Pembahasan sigma_s dan sigma_r)
    p = create_p()
    format_p(p)
    add_styled_run(p, "Kinerja penapisan Bilateral Filter dikendalikan secara presisi oleh interaksi dua parameter independen, yaitu simpangan baku spasial (")
    add_styled_run(p, "\u03c3s", italic=True)
    add_styled_run(p, ") dan simpangan baku intensitas (")
    add_styled_run(p, "\u03c3r", italic=True)
    add_styled_run(p, "):")

    # Item a
    p = create_p()
    format_p(p, first_line_indent_pt=0)
    p.paragraph_format.left_indent = Pt(28.35)
    add_styled_run(p, "a.\tParameter Spasial (", bold=True)
    add_styled_run(p, "\u03c3s", bold=True, italic=True)
    add_styled_run(p, "): Mengendalikan bobot kedekatan geometris di dalam domain spasial melalui fungsi pembobot ", bold=True)
    add_styled_run(p, "G\u03c3s(||p \u2212 q||) = exp(\u2212||p \u2212 q||\u00b2 / (2\u03c3s\u00b2))", italic=True)
    add_styled_run(p, ". Parameter ini menentukan radius lingkungan piksel tetangga yang dilibatkan dalam perhitungan rata-rata terbobot. Nilai ")
    add_styled_run(p, "\u03c3s", italic=True)
    add_styled_run(p, " yang moderat diperlukan pada radiografi toraks untuk menangkap korelasi spasial fluktuasi derau kuantum ")
    add_styled_run(p, "mottle", italic=True)
    add_styled_run(p, ", namun nilai yang terlalu besar berpotensi mengaburkan detail kurvatur anatomis lokal yang halus.")

    # Item b
    p = create_p()
    format_p(p, first_line_indent_pt=0)
    p.paragraph_format.left_indent = Pt(28.35)
    add_styled_run(p, "b.\tParameter Rentang Intensitas (", bold=True)
    add_styled_run(p, "\u03c3r", bold=True, italic=True)
    add_styled_run(p, "): Mengendalikan sensitivitas filter terhadap variasi gradien intensitas radiometrik melalui fungsi pembobot ", bold=True)
    add_styled_run(p, "G\u03c3r(|I(p) \u2212 I(q)|) = exp(\u2212|I(p) \u2212 I(q)|\u00b2 / (2\u03c3r\u00b2))", italic=True)
    add_styled_run(p, ". Parameter ")
    add_styled_run(p, "\u03c3r", italic=True)
    add_styled_run(p, " berperan sebagai gerbang selektivitas ambang batas fotometrik. Pada area jaringan homogen di mana selisih |")
    add_styled_run(p, "I(p) \u2212 I(q)", italic=True)
    add_styled_run(p, "| kecil, bobot rentang bernilai mendekati satu (")
    add_styled_run(p, "G\u03c3r \u2248 1", italic=True)
    add_styled_run(p, ") sehingga filter bekerja meratakan derau. Sebaliknya, pada area garis batas tepi kardiak di mana selisih intensitas besar, bobot rentang seketika anjlok mendekati nol (")
    add_styled_run(p, "G\u03c3r \u2192 0", italic=True)
    add_styled_run(p, "). Akibatnya, piksel dari sisi seberang garis batas diabaikan dalam pembobotan, menjaga batas terluar kardiak tetap terlindungi dari efek pengaburan (")
    add_styled_run(p, "edge preservation", italic=True)
    add_styled_run(p, ").")

    # Paragraf Pengantar Gambar 2.2
    p = create_p()
    format_p(p)
    add_styled_run(p, "Prinsip kerja penapisan bilateral yang memadukan bobot kedekatan geometris domain spasial dan bobot kemiripan radiometrik dalam mempertahankan ketajaman batas tepi kardiak diilustrasikan secara komprehensif pada Gambar 2.2.")

    # Paragraf Gambar 2.2 (Preserved drawing)
    p_img_22 = create_p()
    format_p(p_img_22, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, line_spacing=1.0, first_line_indent_pt=0)
    if drawing_fig2_2 is not None:
        run_img = p_img_22.add_run()
        run_img._r.append(drawing_fig2_2)

    # Caption Gambar 2.2
    p_cap_22 = create_p()
    format_p(p_cap_22, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=12, line_spacing=1.0, first_line_indent_pt=0)
    add_styled_run(p_cap_22, "Gambar 2. 2 Prinsip Kerja Bilateral Filtering dalam Mereduksi Derau sembari Mempertahankan Batas Tepi Organ (Tomasi & Manduchi, 1998)", size_pt=11)

    # =========================================================================
    # HEADING 3: 2.4.3 Denoising Convolutional Neural Network (DnCNN)
    # =========================================================================
    h3_3 = create_p()
    format_p(h3_3, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=4, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h3_3, "2.4.3\tDenoising Convolutional Neural Network (DnCNN)", bold=True, size_pt=12)

    # Paragraf 1 DnCNN
    p = create_p()
    format_p(p)
    add_styled_run(p, "Denoising Convolutional Neural Network (DnCNN) merupakan arsitektur jaringan saraf tiruan konvolusional dalam (")
    add_styled_run(p, "deep CNN", italic=True)
    add_styled_run(p, ") yang memanfaatkan integrasi teknik ")
    add_styled_run(p, "residual learning", italic=True)
    add_styled_run(p, " dan ")
    add_styled_run(p, "Batch Normalization", italic=True)
    add_styled_run(p, " (BN) untuk mengatasi permasalahan restorasi citra berderau secara komprehensif (Zhang et al., 2017). Berbeda dengan paradigma model CNN konvensional yang dilatih secara langsung untuk memetakan citra kotor ")
    add_styled_run(p, "y", italic=True)
    add_styled_run(p, " menjadi estimasi citra bersih ")
    add_styled_run(p, "x\u0302", italic=True)
    add_styled_run(p, ", strategi pembelajaran residual pada DnCNN dirancang untuk memprediksi matriks derau residual yang terkandung di dalam citra, dinotasikan sebagai ")
    add_styled_run(p, "R(y) \u2248 v", italic=True)
    add_styled_run(p, ", di mana ")
    add_styled_run(p, "y = x + v", italic=True)
    add_styled_run(p, ". Dengan mengalihkan fokus estimasi jaringan ke residual derau, citra bersih hasil restorasi diperoleh melalui operasi pengurangan elementer:")

    # Eq 5: x_hat = y - R(y)
    create_eq_p('<math xmlns="http://www.w3.org/1998/Math/MathML"><mover><mi>x</mi><mo>^</mo></mover><mo>=</mo><mi>y</mi><mo>-</mo><mi mathvariant="script">R</mi><mo>(</mo><mi>y</mi><mo>)</mo></math>')

    # Paragraf 2 DnCNN (Loss Function)
    p = create_p()
    format_p(p)
    add_styled_run(p, "Strategi pembelajaran residual ini mempermudah proses optimasi gradien karena distribusi matriks derau residual ")
    add_styled_run(p, "v", italic=True)
    add_styled_run(p, " jauh lebih terpusat dan mendekati Gaussian acak dibandingkan struktur manifold citra radiografi yang kompleks. Hal ini mencegah permasalahan hilangnya gradien (")
    add_styled_run(p, "vanishing gradient problem", italic=True)
    add_styled_run(p, ") serta mempercepat laju konvergensi pelatihan (He et al., 2016). Fungsi objektif kerugian (")
    add_styled_run(p, "loss function", italic=True)
    add_styled_run(p, ") dilatih meminimalkan ")
    add_styled_run(p, "Mean Squared Error", italic=True)
    add_styled_run(p, " (MSE) residual:")

    # Eq 6: DnCNN MSE Loss
    create_eq_p('<math xmlns="http://www.w3.org/1998/Math/MathML"><mi mathvariant="script">L</mi><mo>(</mo><mi>Θ</mi><mo>)</mo><mo>=</mo><mfrac><mn>1</mn><mrow><mn>2</mn><mi>N</mi></mrow></mfrac><munderover><mo>∑</mo><mrow><mi>i</mi><mo>=</mo><mn>1</mn></mrow><mi>N</mi></munderover><msubsup><mrow><mo>‖</mo><mi mathvariant=\"script\">R</mi><mo>(</mo><msub><mi>y</mi><mi>i</mi></msub><mo>;</mo><mi>Θ</mi><mo>)</mo><mo>-</mo><mo>(</mo><msub><mi>y</mi><mi>i</mi></msub><mo>-</mo><msub><mi>x</mi><mi>i</mi></msub><mo>)</mo><mo>‖</mo></mrow><mi>F</mi><mn>2</mn></msubsup></math>')

    # Paragraf 3 DnCNN (Rincian 17 lapisan)
    p = create_p()
    format_p(p)
    add_styled_run(p, "Arsitektur DnCNN terdiri atas 17 lapisan konvolusional bertingkat tanpa lapisan ")
    add_styled_run(p, "pooling", italic=True)
    add_styled_run(p, " guna menjaga integritas resolusi spasial citra tetap utuh 1:1, dengan rincian konfigurasi sebagai berikut:")

    # Item 1 DnCNN
    p = create_p()
    format_p(p, first_line_indent_pt=0)
    p.paragraph_format.left_indent = Pt(28.35)
    add_styled_run(p, "1.\tLapisan Pertama (Input Layer): ", bold=True)
    add_styled_run(p, "Terdiri atas operasi konvolusi tunggal Conv + ReLU menggunakan 64 filter berukuran 3 \u00d7 3 \u00d7 ")
    add_styled_run(p, "c", italic=True)
    add_styled_run(p, " (di mana ")
    add_styled_run(p, "c = 1", italic=True)
    add_styled_run(p, " untuk citra radiografi monokromatik abu-abu), menghasilkan 64 peta fitur (")
    add_styled_run(p, "feature maps", italic=True)
    add_styled_run(p, ") awal tanpa menerapkan Batch Normalization.")

    # Item 2 DnCNN
    p = create_p()
    format_p(p, first_line_indent_pt=0)
    p.paragraph_format.left_indent = Pt(28.35)
    add_styled_run(p, "2.\tLapisan 2 hingga 16 (Intermediate Processing Layers): ", bold=True)
    add_styled_run(p, "Terdiri atas 15 blok berturut-turut dengan konfigurasi Conv + Batch Normalization (BN) + ReLU (64 filter berukuran 3 \u00d7 3 \u00d7 64). Integrasi Batch Normalization menstabilkan distribusi input internal (")
    add_styled_run(p, "internal covariate shift", italic=True)
    add_styled_run(p, "), mempercepat perambatan gradien, serta bertindak sebagai regularisator implisit. Bidang pandang efektif (")
    add_styled_run(p, "receptive field", italic=True)
    add_styled_run(p, ") membesar secara linier mencapai ukuran 35 \u00d7 35 piksel pada lapisan ke-17.")

    # Item 3 DnCNN
    p = create_p()
    format_p(p, first_line_indent_pt=0)
    p.paragraph_format.left_indent = Pt(28.35)
    add_styled_run(p, "3.\tLapisan Terakhir (Reconstruction Layer): ", bold=True)
    add_styled_run(p, "Terdiri atas satu operasi konvolusi Conv dengan 1 filter berukuran 3 \u00d7 3 \u00d7 64 tanpa aktivasi non-linear maupun BN guna merekonstruksi 64 peta fitur akhir menjadi matriks tunggal estimasi residual derau.")

    # Paragraf Pengantar Gambar 2.3
    p = create_p()
    format_p(p)
    add_styled_run(p, "Konfigurasi 17 lapisan konvolusional bertingkat dan mekanisme pemisahan derau residual pada arsitektur DnCNN disajikan secara sistematis pada Gambar 2.3.")

    # Paragraf Gambar 2.3 (Preserved drawing)
    p_img_23 = create_p()
    format_p(p_img_23, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, line_spacing=1.0, first_line_indent_pt=0)
    if drawing_fig2_3 is not None:
        run_img = p_img_23.add_run()
        run_img._r.append(drawing_fig2_3)

    # Caption Gambar 2.3
    p_cap_23 = create_p()
    format_p(p_cap_23, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=12, line_spacing=1.0, first_line_indent_pt=0)
    add_styled_run(p_cap_23, "Gambar 2. 3 Arsitektur Jaringan DnCNN Berbasis Residual Learning untuk Pemisahan Derau Citra (Zhang et al., 2017)", size_pt=11)

    # =========================================================================
    # HEADING 3: 2.4.4 Domain Shift dan Tantangan Deep Denoising pada Radiografi Medis
    # =========================================================================
    h3_4 = create_p()
    format_p(h3_4, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=4, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h3_4, "2.4.4\tDomain Shift dan Tantangan Deep Denoising pada Radiografi Medis", bold=True, size_pt=12)

    # Paragraf 1 Domain Shift
    p = create_p()
    format_p(p)
    add_styled_run(p, "Meskipun arsitektur DnCNN memiliki kapabilitas restorasi yang sangat tinggi, penerapannya secara langsung pada citra radiografi toraks menghadapi tantangan mendasar yang dikenal sebagai pergeseran domain (")
    add_styled_run(p, "domain shift", italic=True)
    add_styled_run(p, ") atau ketidakcocokan domain (")
    add_styled_run(p, "domain mismatch", italic=True)
    add_styled_run(p, ") (Guan & Liu, 2022). Model DnCNN orisinal dikembangkan dan dilatih menggunakan dataset citra alami fotografis (seperti dataset BSD500) dengan asumsi derau sintetis aditif bebas sinyal (")
    add_styled_run(p, "Additive White Gaussian Noise", italic=True)
    add_styled_run(p, " atau AWGN). Karakteristik statistik citra alami tersebut memiliki perbedaan fisik yang substansial dibandingkan dengan radiografi toraks medis, di mana derau radiasi sinar-X didominasi oleh fluktuasi foton Poisson yang bergantung sinyal (")
    add_styled_run(p, "signal-dependent quantum mottle", italic=True)
    add_styled_run(p, ") dan citra klinis memiliki rentang kedalaman fotometrik tinggi (12-bit atau 16-bit) dengan tumpang tindih struktur anatomi proyeksi yang sangat kompleks.")

    # Paragraf 2 Transfer Learning & Fine-Tuning
    p = create_p()
    format_p(p)
    add_styled_run(p, "Untuk mengatasi kesenjangan domain tersebut, penelitian ini menerapkan pendekatan adaptasi domain melalui strategi ")
    add_styled_run(p, "transfer learning", italic=True)
    add_styled_run(p, " dan penyesuaian bobot (")
    add_styled_run(p, "fine-tuning", italic=True)
    add_styled_run(p, "). Bobot awal model dasar (")
    add_styled_run(p, "pretrained weights", italic=True)
    add_styled_run(p, ") dieksploitasi untuk mendeteksi fitur spasial primitif (seperti detektor tepi dan tekstur lokal), yang kemudian dilatih ulang secara terarah menggunakan dataset radiografi toraks digital (seperti JSRT dan VinDr-CXR) yang disuntikkan model derau campuran Poisson-Gaussian terkontrol.")

    # Paragraf 3 Tiga Risiko Metodologis
    p = create_p()
    format_p(p)
    add_styled_run(p, "Dalam perancangan dan evaluasi restorasi berbasis ")
    add_styled_run(p, "deep learning", italic=True)
    add_styled_run(p, " untuk citra medis, terdapat tiga risiko metodologis kritis yang wajib diantisipasi dan dimitigasi:")

    # Poin a: Data Leakage
    p = create_p()
    format_p(p, first_line_indent_pt=0)
    p.paragraph_format.left_indent = Pt(28.35)
    add_styled_run(p, "a.\tPencegahan Kebocoran Data (", bold=True)
    add_styled_run(p, "Data Leakage Prevention", bold=True, italic=True)
    add_styled_run(p, "): Pembagian dataset pelatihan, validasi, dan pengujian wajib dilakukan secara ketat berbasis identitas pasien (", bold=True)
    add_styled_run(p, "patient-wise split", italic=True)
    add_styled_run(p, "), bukan pemisahan acak berbasis citra atau ")
    add_styled_run(p, "patch", italic=True)
    add_styled_run(p, ". Hal ini krusial guna memastikan model tidak mengevaluasi citra dari pasien yang anatomisnya telah dipelajari pada fase pelatihan, yang dapat menghasilkan evaluasi performa yang overestimasi secara semu (")
    add_styled_run(p, "optimistic bias", italic=True)
    add_styled_run(p, ").")

    # Poin b: Over-smoothing
    p = create_p()
    format_p(p, first_line_indent_pt=0)
    p.paragraph_format.left_indent = Pt(28.35)
    add_styled_run(p, "b.\tRisiko Penghalusan Berlebih (", bold=True)
    add_styled_run(p, "Over-Smoothing Risk", bold=True, italic=True)
    add_styled_run(p, "): Konvolusi dalam berisiko menginterpretasikan struktur mikro anatomi halus berkontras rendah (seperti corakan vaskular perifer paru atau kontur tipis kardioperikardial) sebagai komponen derau, sehingga struktur tersebut terhapus secara berlebihan. Pada konteks pengukuran CTR, penghalusan berlebih pada apeks kardiak lateral dapat mengikis batas ventrikel kiri, menyebabkan diameter jantung terukur menyempit secara artifisial.")

    # Poin c: Hallucination Risk
    p = create_p()
    format_p(p, first_line_indent_pt=0)
    p.paragraph_format.left_indent = Pt(28.35)
    add_styled_run(p, "c.\tRisiko Halusinasi Rekonstruksi (", bold=True)
    add_styled_run(p, "Hallucination Risk", bold=True, italic=True)
    add_styled_run(p, "): Sebagai jaringan saraf non-linear dalam, model memiliki risiko memunculkan pola sintetis palsu (", bold=True)
    add_styled_run(p, "hallucination artifacts", italic=True)
    add_styled_run(p, ") yang tidak bersumber dari transmisi radiasi fisik asli (Bhadra et al., 2021). Pada radiografi toraks, halusinasi dapat mengubah kurvatur batas lateral kardiak atau membentuk garis semu yang menyerupai infiltrat paru, yang berpotensi fatal bagi akurasi pengukuran diagnostik biometri. Oleh karena itu, pengujian performa DnCNN dalam penelitian ini tidak semata dievaluasi melalui metrik kualitas umum (PSNR dan SSIM), melainkan wajib divalidasi langsung melalui metrik berbasis tugas (")
    add_styled_run(p, "task-based evaluation", italic=True)
    add_styled_run(p, ") terhadap galat pengukuran CTR dan pelestarian batas kontur anatomi (")
    add_styled_run(p, "boundary preservation", italic=True)
    add_styled_run(p, ").")

    # =========================================================================
    # DELETE OLD SECTION 2.4 PARAGRAPHS
    # =========================================================================
    print(f"Removing old paragraphs from index {start_idx} to {end_idx - 1}...")
    for p_old in doc.paragraphs[start_idx:end_idx]:
        p_old._p.getparent().remove(p_old._p)

    doc.save(doc_path)
    print(f"Successfully updated document {doc_path}!")
    print(f"Total final paragraphs: {len(doc.paragraphs)}")

if __name__ == '__main__':
    run()
