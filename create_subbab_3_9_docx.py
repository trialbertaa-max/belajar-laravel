import os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import lxml.etree as ET

xslt_path = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
xslt = ET.parse(xslt_path)
transform = ET.XSLT(xslt)

def mml_to_omml_element(mml_str):
    mml_tree = ET.fromstring(mml_str)
    omml_tree = transform(mml_tree)
    root = omml_tree.getroot()
    if root.tag.endswith('oMath'):
        para = ET.Element('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara')
        para_pr = ET.SubElement(para, '{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathParaPr')
        jc = ET.SubElement(para_pr, '{http://schemas.openxmlformats.org/officeDocument/2006/math}jc')
        jc.set('{http://schemas.openxmlformats.org/officeDocument/2006/math}val', 'center')
        para.append(root)
        return para
    return root

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False, subscript=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic
    r.font.subscript = subscript

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, space_before_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.space_before = Pt(space_before_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, size_pt=12, subscript=False):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, size_pt=size_pt, subscript=subscript)
    return r

def main():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.18)    # 3 cm
        s.bottom_margin = Inches(1.18) # 3 cm
        s.left_margin = Inches(1.57)   # 4 cm
        s.right_margin = Inches(1.18)  # 3 cm

    # Heading 3.9
    h = doc.add_paragraph()
    format_p(h, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=6, space_before_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h, '3.9\tProsedur Pengukuran CTR dan Kaidah Landmark', bold=True)

    # Paragraf 1: Definisi & Pengantar Rumus Danzer
    p1 = doc.add_paragraph()
    format_p(p1)
    add_styled_run(p1, 'Pengukuran ')
    add_styled_run(p1, 'Cardiothoracic Ratio', italic=True)
    add_styled_run(p1, ' (CTR) dilakukan mengikuti kaidah standar Danzer (1919) untuk mengukur rasio antara diameter transversal jantung dan sangkar toraks sebagai berikut:')

    # Persamaan CTR Danzer (Rata tengah tanpa nomor persamaan)
    mml_ctr = '''<math xmlns="http://www.w3.org/1998/Math/MathML">
      <mrow>
        <mtext>CTR</mtext>
        <mo>=</mo>
        <mfrac>
          <mrow>
            <mtext>MRD</mtext>
            <mo>+</mo>
            <mtext>MLD</mtext>
          </mrow>
          <mtext>MTD</mtext>
        </mfrac>
        <mo>&#215;</mo>
        <mn>100</mn>
        <mo>%</mo>
      </mrow>
    </math>'''

    p_eq = doc.add_paragraph()
    format_p(p_eq, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, space_before_pt=4, line_spacing=1.0, first_line_indent_pt=0)
    omml_ctr = mml_to_omml_element(mml_ctr)
    p_eq._p.append(omml_ctr)

    # Paragraf 2: Keterangan Parameter & Pixel Spacing
    p2 = doc.add_paragraph()
    format_p(p2)
    add_styled_run(p2, 'dengan MRD (')
    add_styled_run(p2, 'Maximum Right Diameter', italic=True)
    add_styled_run(p2, ') dan MLD (')
    add_styled_run(p2, 'Maximum Left Diameter', italic=True)
    add_styled_run(p2, ') adalah proyeksi jarak terjauh batas jantung kanan dan kiri ke garis sumbu tengah, serta MTD (')
    add_styled_run(p2, 'Maximum Thoracic Diameter', italic=True)
    add_styled_run(p2, ') adalah diameter transversal toraks dalam. Sebagai rasio fraksional, kalkulasi CTR valid dilakukan langsung pada domain piksel maupun terkalibrasi via ')
    add_styled_run(p2, 'Pixel Spacing', italic=True)
    add_styled_run(p2, ' (0028, 0030).')

    # Paragraf 3: Pengantar Kaidah Landmark
    p3 = doc.add_paragraph()
    format_p(p3)
    add_styled_run(p3, 'Penentuan titik batas ukur (')
    add_styled_run(p3, 'landmarks', italic=True)
    add_styled_run(p3, ') diatur secara kaku berdasarkan kaidah anatomis:')

    # 4 Kaidah Landmark Ringkas
    landmarks = [
        ("1.", "Garis sumbu tengah (midline): Garis vertikal melalui titik tengah prosesus spinosus vertebralis torakalis."),
        ("2.", "Batas kardiak kanan (MRD): Garis horizontal tegak lurus midline menuju batas lateral atrium kanan."),
        ("3.", "Batas kardiak kiri (MLD): Garis horizontal tegak lurus midline menuju apeks ventrikel kiri."),
        ("4.", "Diameter toraks (MTD): Garis horizontal pada puncak hemidiafragma kanan dari tepi dalam kosta kanan ke kosta kiri.")
    ]

    for num, text in landmarks:
        p_lm = doc.add_paragraph()
        format_p(p_lm, space_after_pt=1, space_before_pt=0, first_line_indent_pt=14.17)
        add_styled_run(p_lm, num + '\t')
        words = text.split(' ')
        for w_idx, w in enumerate(words):
            sp = ' ' if w_idx < len(words) - 1 else ''
            clean_w = w.strip('(),:.')
            is_it = clean_w in ['midline', 'hemidiafragma', 'prosesus', 'spinosus', 'vertebralis', 'torakalis', 'MRD', 'MLD', 'MTD']
            add_styled_run(p_lm, w + sp, italic=is_it)

    # Paragraf 4: Independent re-detection Ringkas
    p4 = doc.add_paragraph()
    format_p(p4)
    add_styled_run(p4, 'Titik batas anatomi dideteksi ulang secara independen (')
    add_styled_run(p4, 'independent re-detection', italic=True)
    add_styled_run(p4, ') pada setiap kondisi citra menggunakan model segmentasi yang sama guna menguji pergeseran ')
    add_styled_run(p4, 'landmark', italic=True)
    add_styled_run(p4, ' akibat perlakuan restorasi secara objektif.')

    # Gambar 3.3
    img_path = r'd:\laragon\www\belajar-laravel\Gambar_3_3_Skema_CTR_Danzer.png'
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        format_p(p_img, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, space_before_pt=8, line_spacing=1.0, first_line_indent_pt=0)
        p_img.add_run().add_picture(img_path, width=Inches(5.0))

        # Caption Gambar 3.3
        p_cap = doc.add_paragraph()
        format_p(p_cap, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=6, space_before_pt=2, line_spacing=1.15, first_line_indent_pt=0)
        add_styled_run(p_cap, 'Gambar 3.3 Skema Penentuan Region of Interest dan Penarikan Garis Ukur CTR Metode Danzer (1919)', bold=False)

    out_file = r'd:\laragon\www\belajar-laravel\SUBBAB_3_9_PENGUKURAN_CTR.docx'
    doc.save(out_file)
    print('Successfully updated Subbab 3.9 (Ultra-Ringkas):', out_file)

if __name__ == '__main__':
    main()
