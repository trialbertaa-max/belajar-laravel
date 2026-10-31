import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.text.paragraph import Paragraph
import lxml.etree as ET
import shutil
import copy

# Load XSLT
xslt_path = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
xslt = ET.parse(xslt_path)
transform = ET.XSLT(xslt)

def mml_to_omml_element(mml_str):
    mml_tree = ET.fromstring(mml_str)
    omml_tree = transform(mml_tree)
    root = omml_tree.getroot()
    if root.tag.endswith('oMath'):
        para = ET.Element('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara')
        para_pr = ET.SubElement(para, '{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathParaPr')
        jc = ET.SubElement(para_pr, '{http://schemas.openxmlformats.org/officeDocument/2006/math}jc')
        jc.set('{http://schemas.openxmlformats.org/officeDocument/2006/math}val', 'center')
        para.append(root)
        root = para
    xml_bytes = ET.tostring(root, encoding='utf-8')
    return parse_xml(xml_bytes)

def set_run_font(r, font_name='Times New Roman', size_pt=12, bold=False, italic=False, subscript=False, superscript=False):
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    r.italic = italic
    if subscript:
        r.font.subscript = True
    if superscript:
        r.font.superscript = True

def format_p(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after_pt=0, line_spacing=1.5, first_line_indent_pt=28.35):
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after_pt)
    p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)

def add_styled_run(p, text, bold=False, italic=False, subscript=False, superscript=False, size_pt=12):
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic, subscript=subscript, superscript=superscript, size_pt=size_pt)
    return r

def build_section_243(target_before_p, drawing_fig2_3):
    def create_p(style_name='Normal'):
        new_p_elm = OxmlElement('w:p')
        target_before_p._p.addprevious(new_p_elm)
        p = Paragraph(new_p_elm, target_before_p._parent)
        if style_name != 'Normal':
            p.style = style_name
        return p

    def create_eq_p(mml_str):
        p = create_p()
        format_p(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, line_spacing=1.0, first_line_indent_pt=0)
        omml_elm = mml_to_omml_element(mml_str)
        p._p.append(omml_elm)
        return p

    # 1. Heading 3: 2.4.3 Denoising Convolutional Neural Network (DnCNN)
    h3 = create_p('Heading 3')
    format_p(h3, align=WD_ALIGN_PARAGRAPH.LEFT, space_after_pt=4, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(h3, "2.4.3\t", bold=True, italic=False)
    add_styled_run(h3, "Denoising Convolutional Neural Network", bold=True, italic=True)
    add_styled_run(h3, " (DnCNN)", bold=True, italic=False)

    # 2. Paragraf 1: Paradigma Arsitektur & Residual Learning
    p1 = create_p()
    format_p(p1)
    add_styled_run(p1, "Denoising Convolutional Neural Network", italic=True)
    add_styled_run(p1, " (DnCNN) merupakan arsitektur jaringan saraf tiruan konvolusional dalam (")
    add_styled_run(p1, "deep convolutional neural network", italic=True)
    add_styled_run(p1, ") yang dirancang secara khusus untuk memisahkan derau dari sinyal laten citra asli melalui integrasi strategi pembelajaran ")
    add_styled_run(p1, "residual", italic=True)
    add_styled_run(p1, " (")
    add_styled_run(p1, "residual learning", italic=True)
    add_styled_run(p1, ") dan penormalan tumpak (")
    add_styled_run(p1, "Batch Normalization", italic=True)
    add_styled_run(p1, ") (Zhang et al., 2017). Berbeda dengan paradigma model restorasi konvensional berbasis pembelajaran mendalam (")
    add_styled_run(p1, "deep learning", italic=True)
    add_styled_run(p1, ") yang dilatih memetakan citra berderau ")
    add_styled_run(p1, "y", italic=True)
    add_styled_run(p1, " secara langsung menuju estimasi citra bersih ")
    add_styled_run(p1, "x̂", italic=True)
    add_styled_run(p1, ", strategi pembelajaran ")
    add_styled_run(p1, "residual", italic=True)
    add_styled_run(p1, " pada DnCNN dirancang untuk memprediksi matriks sisa derau (")
    add_styled_run(p1, "noise residual", italic=True)
    add_styled_run(p1, ") yang mengontaminasi citra, dinotasikan sebagai ")
    add_styled_run(p1, "R", italic=True)
    add_styled_run(p1, "(")
    add_styled_run(p1, "y", italic=True)
    add_styled_run(p1, ") \u2248 ")
    add_styled_run(p1, "v", italic=True)
    add_styled_run(p1, ", dengan model pembentukan degradasi aditif ")
    add_styled_run(p1, "y", italic=True)
    add_styled_run(p1, " = ")
    add_styled_run(p1, "x", italic=True)
    add_styled_run(p1, " + ")
    add_styled_run(p1, "v", italic=True)
    add_styled_run(p1, ". Melalui formulasi ini, estimasi citra anatomis bersih ")
    add_styled_run(p1, "x̂", italic=True)
    add_styled_run(p1, " diperoleh melalui operasi pengurangan elementer sederhana antara citra masukan berderau dan keluaran sisa derau yang diprediksi oleh jaringan:")

    # 3. Persamaan 1: x_hat = y - R(y)
    mml1 = '''<math xmlns="http://www.w3.org/1998/Math/MathML">
      <mover><mi>x</mi><mo>^</mo></mover>
      <mo>=</mo>
      <mi>y</mi>
      <mo>-</mo>
      <mi mathvariant="script">R</mi>
      <mo>(</mo><mi>y</mi><mo>)</mo>
    </math>'''
    create_eq_p(mml1)

    # 4. Paragraf 2: Keunggulan Konvergensi & Formulasi Loss Function
    p2 = create_p()
    format_p(p2)
    add_styled_run(p2, "Pengalihan target estimasi menuju pemodelan sisa derau memberikan keunggulan komputasional yang mendasar dalam proses optimasi gradien. Distribusi statistik sisa derau ")
    add_styled_run(p2, "v", italic=True)
    add_styled_run(p2, " memiliki manifold geometris yang jauh lebih sederhana, terpusat, dan mendekati distribusi ")
    add_styled_run(p2, "Gaussian", italic=True)
    add_styled_run(p2, " acak dibandingkan manifold struktur fitur anatomis citra radiografi yang sangat kompleks dan berdimensi tinggi. Karakteristik ini secara efektif mencegah timbulnya degradasi performa akibat hilangnya gradien (")
    add_styled_run(p2, "vanishing gradient problem", italic=True)
    add_styled_run(p2, ") pada jaringan yang dalam serta mempercepat laju konvergensi proses pelatihan (He et al., 2016). Selama proses pelatihan, parameter bobot dan bias jaringan ")
    add_styled_run(p2, "Θ", italic=True)
    add_styled_run(p2, " dioptimalkan secara terawasi (")
    add_styled_run(p2, "supervised learning", italic=True)
    add_styled_run(p2, ") menggunakan kumpulan pasangan data citra radiografi berderau dan bersih (")
    add_styled_run(p2, "y", italic=True)
    add_styled_run(p2, "i", italic=True, subscript=True)
    add_styled_run(p2, ", ")
    add_styled_run(p2, "x", italic=True)
    add_styled_run(p2, "i", italic=True, subscript=True)
    add_styled_run(p2, ") dengan meminimalkan fungsi kerugian (")
    add_styled_run(p2, "loss function", italic=True)
    add_styled_run(p2, ") berbasis galat kuadrat rata-rata (")
    add_styled_run(p2, "Mean Squared Error", italic=True)
    add_styled_run(p2, " atau MSE) ")
    add_styled_run(p2, "residual", italic=True)
    add_styled_run(p2, ":")

    # 5. Persamaan 2: Loss Function MSE Residual
    mml2 = '''<math xmlns="http://www.w3.org/1998/Math/MathML">
      <mi mathvariant="script">L</mi>
      <mo>(</mo><mi>Θ</mi><mo>)</mo>
      <mo>=</mo>
      <mfrac>
        <mn>1</mn>
        <mrow><mn>2</mn><mi>N</mi></mrow>
      </mfrac>
      <munderover>
        <mo>∑</mo>
        <mrow><mi>i</mi><mo>=</mo><mn>1</mn></mrow>
        <mi>N</mi>
      </munderover>
      <msubsup>
        <mrow>
          <mo>‖</mo>
          <mi mathvariant="script">R</mi>
          <mo>(</mo><msub><mi>y</mi><mi>i</mi></msub><mo>;</mo><mi>Θ</mi><mo>)</mo>
          <mo>-</mo>
          <mo>(</mo><msub><mi>y</mi><mi>i</mi></msub><mo>-</mo><msub><mi>x</mi><mi>i</mi></msub><mo>)</mo>
          <mo>‖</mo>
        </mrow>
        <mi>F</mi>
        <mn>2</mn>
      </msubsup>
    </math>'''
    create_eq_p(mml2)

    p2_desc = create_p()
    format_p(p2_desc)
    add_styled_run(p2_desc, "di mana ")
    add_styled_run(p2_desc, "N", italic=True)
    add_styled_run(p2_desc, " merepresentasikan jumlah cuplikan bidang citra (")
    add_styled_run(p2_desc, "image patches", italic=True)
    add_styled_run(p2_desc, "), ")
    add_styled_run(p2_desc, "R", italic=True)
    add_styled_run(p2_desc, "(")
    add_styled_run(p2_desc, "y", italic=True)
    add_styled_run(p2_desc, "i", italic=True, subscript=True)
    add_styled_run(p2_desc, "; ")
    add_styled_run(p2_desc, "Θ", italic=True)
    add_styled_run(p2_desc, ") adalah estimasi sisa derau yang dihasilkan oleh model jaringan, (")
    add_styled_run(p2_desc, "y", italic=True)
    add_styled_run(p2_desc, "i", italic=True, subscript=True)
    add_styled_run(p2_desc, " \u2212 ")
    add_styled_run(p2_desc, "x", italic=True)
    add_styled_run(p2_desc, "i", italic=True, subscript=True)
    add_styled_run(p2_desc, ") adalah sisa derau acuan (")
    add_styled_run(p2_desc, "ground truth noise residual", italic=True)
    add_styled_run(p2_desc, "), dan ||\u00b7||")
    add_styled_run(p2_desc, "F", italic=True, subscript=True)
    add_styled_run(p2_desc, " melambangkan norma Frobenius matriks.")

    # 6. Paragraf 3: Konfigurasi 17 Lapisan
    p3 = create_p()
    format_p(p3)
    add_styled_run(p3, "Arsitektur standar DnCNN yang diterapkan dalam penelitian ini tersusun atas 17 lapisan konvolusional bertingkat tanpa menerapkan lapisan penggabungan (")
    add_styled_run(p3, "pooling layer", italic=True)
    add_styled_run(p3, ") maupun lapisan terhubung penuh (")
    add_styled_run(p3, "fully-connected layer", italic=True)
    add_styled_run(p3, "). Peniadaan operasi ")
    add_styled_run(p3, "pooling", italic=True)
    add_styled_run(p3, " merupakan aspek desain krusial untuk mencegah penurunan resolusi spasial citra (")
    add_styled_run(p3, "downsampling", italic=True)
    add_styled_run(p3, "), sehingga integritas ukuran matriks dan ketajaman batas mikrostruktur anatomis toraks tetap terpelihara pada skala resolusi asli 1:1. Konfigurasi 17 lapisan tersebut terbagi ke dalam tiga tingkatan fungsional:")

    # Poin 1: Lapisan Masukan
    p_layer1 = create_p()
    format_p(p_layer1)
    add_styled_run(p_layer1, "1.\t", bold=True, italic=False)
    add_styled_run(p_layer1, "Lapisan Masukan (", bold=True, italic=False)
    add_styled_run(p_layer1, "Input Layer", bold=True, italic=True)
    add_styled_run(p_layer1, "): ", bold=True, italic=False)
    add_styled_run(p_layer1, "Terdiri atas satu operasi konvolusi tunggal Conv + ReLU (")
    add_styled_run(p_layer1, "Rectified Linear Unit", italic=True)
    add_styled_run(p_layer1, ") yang menggunakan 64 penapis (")
    add_styled_run(p_layer1, "filters", italic=True)
    add_styled_run(p_layer1, ") berukuran 3 \u00d7 3 \u00d7 ")
    add_styled_run(p_layer1, "c", italic=True)
    add_styled_run(p_layer1, " (dengan ")
    add_styled_run(p_layer1, "c = 1", italic=True)
    add_styled_run(p_layer1, " untuk citra radiografi digital monokromatik abu-abu berskala 12-bit atau 16-bit). Lapisan awal ini mengekstraksi 64 peta fitur spasial (")
    add_styled_run(p_layer1, "feature maps", italic=True)
    add_styled_run(p_layer1, ") dari citra masukan tanpa menerapkan ")
    add_styled_run(p_layer1, "Batch Normalization", italic=True)
    add_styled_run(p_layer1, ".")

    # Poin 2: Lapisan Pengolah Antara
    p_layer2 = create_p()
    format_p(p_layer2)
    add_styled_run(p_layer2, "2.\t", bold=True, italic=False)
    add_styled_run(p_layer2, "Lapisan Pengolah Antara (", bold=True, italic=False)
    add_styled_run(p_layer2, "Intermediate Processing Layers", bold=True, italic=True)
    add_styled_run(p_layer2, "): ", bold=True, italic=False)
    add_styled_run(p_layer2, "Terdiri atas 15 blok konvolusi berturut-turut (lapisan ke-2 hingga ke-16) dengan konfigurasi berulang Conv + BN + ReLU. Masing-masing lapisan menerapkan 64 penapis berukuran 3 \u00d7 3 \u00d7 64. Integrasi ")
    add_styled_run(p_layer2, "Batch Normalization", italic=True)
    add_styled_run(p_layer2, " (BN) di antara konvolusi dan aktivasi non-linear berfungsi menstabilkan distribusi masukan internal (")
    add_styled_run(p_layer2, "internal covariate shift", italic=True)
    add_styled_run(p_layer2, "), mereduksi ketergantungan model terhadap inisialisasi parameter awal, mempercepat perambatan gradien, serta bertindak sebagai regularisator implisit. Penumpukan 17 lapisan konvolusi dengan ukuran ")
    add_styled_run(p_layer2, "kernel", italic=True)
    add_styled_run(p_layer2, " 3 \u00d7 3 secara bertahap memperluas bidang pandang reseptif efektif (")
    add_styled_run(p_layer2, "effective receptive field", italic=True)
    add_styled_run(p_layer2, ") jaringan hingga mencapai ukuran 35 \u00d7 35 piksel pada lapisan akhir, yang memungkinkan jaringan menangkap korelasi konteks spasial derau pada area sangkar toraks yang lebih luas.")

    # Poin 3: Lapisan Rekonstruksi Akhir
    p_layer3 = create_p()
    format_p(p_layer3)
    add_styled_run(p_layer3, "3.\t", bold=True, italic=False)
    add_styled_run(p_layer3, "Lapisan Rekonstruksi Akhir (", bold=True, italic=False)
    add_styled_run(p_layer3, "Reconstruction Layer", bold=True, italic=True)
    add_styled_run(p_layer3, "): ", bold=True, italic=False)
    add_styled_run(p_layer3, "Terdiri atas satu lapisan konvolusi tunggal Conv menggunakan 1 penapis berukuran 3 \u00d7 3 \u00d7 64 tanpa aktivasi non-linear maupun BN. Lapisan terminal ini mengagregasi 64 peta fitur antara menjadi satu matriks keluaran tunggal yang merepresentasikan prediksi sisa derau ")
    add_styled_run(p_layer3, "R", italic=True)
    add_styled_run(p_layer3, "(")
    add_styled_run(p_layer3, "y", italic=True)
    add_styled_run(p_layer3, ").")

    # 7. Paragraf 4: Pengantar Gambar 2.3
    p4 = create_p()
    format_p(p4)
    add_styled_run(p4, "Mekanisme pemisahan sisa derau dan konfigurasi hierarkis 17 lapisan konvolusional pada arsitektur DnCNN diilustrasikan secara sistematis pada Gambar 2.3.")

    # 8. Gambar 2.3 Paragraph
    p_fig = create_p()
    format_p(p_fig, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=4, line_spacing=1.0, first_line_indent_pt=0)
    drawing_copy = copy.deepcopy(drawing_fig2_3)
    run_fig = p_fig.add_run()
    run_fig._r.append(drawing_copy)

    # 9. Caption Gambar 2.3 Paragraph
    p_cap = create_p('Caption')
    format_p(p_cap, align=WD_ALIGN_PARAGRAPH.CENTER, space_after_pt=12, line_spacing=1.5, first_line_indent_pt=0)
    add_styled_run(p_cap, "Gambar 2. 3 Arsitektur Jaringan DnCNN Berbasis ", bold=False, italic=False)
    add_styled_run(p_cap, "Residual Learning", bold=False, italic=True)
    add_styled_run(p_cap, " untuk Pemisahan Derau Citra (Zhang et al., 2017)", bold=False, italic=False)

    # 10. Paragraf 5: Jembatan Menuju Domain Shift (2.4.4)
    p5 = create_p()
    format_p(p5)
    add_styled_run(p5, "Meskipun arsitektur DnCNN menunjukkan performa pereduksian derau yang sangat superior pada ranah citra alami (")
    add_styled_run(p5, "natural images", italic=True)
    add_styled_run(p5, "), penerapannya secara langsung pada citra radiografi toraks proyeksi PA menghadapi tantangan mendasar berupa pergeseran ranah (")
    add_styled_run(p5, "domain shift", italic=True)
    add_styled_run(p5, "). Arsitektur DnCNN asli dilatih menggunakan citra alami berdimensi 8-bit dengan distribusi derau aditif ")
    add_styled_run(p5, "Gaussian", italic=True)
    add_styled_run(p5, " (")
    add_styled_run(p5, "Additive White Gaussian Noise", italic=True)
    add_styled_run(p5, " atau AWGN), sedangkan citra radiografi digital tersusun atas kedalaman bit tinggi (12-bit hingga 16-bit) dengan derau campuran kuantum Poisson-")
    add_styled_run(p5, "Gaussian", italic=True)
    add_styled_run(p5, " serta memiliki batasan klinis yang ketat terhadap risiko perubahan struktur anatomi jantung dan paru. Fenomena ketidakcocokan ranah (")
    add_styled_run(p5, "domain mismatch", italic=True)
    add_styled_run(p5, ") ini berisiko memunculkan artefak distorsi struktural atau halusinasi visual (")
    add_styled_run(p5, "hallucination risk", italic=True)
    add_styled_run(p5, ") yang dapat menggeser penentuan garis batas kardioperikardial dan mengacaukan penghitungan rasio CTR. Oleh karena itu, diperlukan strategi adaptasi ranah dan pembelajaran alih (")
    add_styled_run(p5, "domain adaptation and transfer learning", italic=True)
    add_styled_run(p5, ") yang komprehensif, sebagaimana dibahas secara mendalam pada Subbab 2.4.4.")

def main():
    doc_path = r'd:\laragon\www\belajar-laravel\DRAFT PROPOSAL SKRIPSI_REVISI_TERBARU_BACKUP.docx'
    test_doc_path = r'd:\laragon\www\belajar-laravel\test_243_result.docx'

    doc = docx.Document(doc_path)
    start_idx = None
    end_idx = None
    drawing_fig2_3 = None

    for i, p in enumerate(doc.paragraphs):
        if '2.4.3' in p.text:
            start_idx = i
        elif 'Gambar 2. 3' in p.text:
            prev_p = doc.paragraphs[i-1]
            drawings = prev_p._p.xpath('.//w:drawing')
            if drawings:
                drawing_fig2_3 = copy.deepcopy(drawings[0])
        elif '2.4.4' in p.text:
            end_idx = i
            break

    print(f"Original 2.4.3 range: p[{start_idx}] to p[{end_idx}]")
    print(f"Extracted drawing_fig2_3: {drawing_fig2_3 is not None}")

    target_before = doc.paragraphs[end_idx]

    # Insert new 2.4.3 paragraphs
    build_section_243(target_before, drawing_fig2_3)

    # Remove old 2.4.3 paragraphs
    old_p_elms = [doc.paragraphs[k]._p for k in range(start_idx, end_idx)]
    for elm in old_p_elms:
        elm.getparent().remove(elm)

    doc.save(test_doc_path)
    print(f"Saved test document to {test_doc_path}")

if __name__ == '__main__':
    main()
